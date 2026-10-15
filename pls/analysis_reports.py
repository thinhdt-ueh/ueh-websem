"""Word reports for the Power Analysis and Sample Size Sensitivity pages.

Both are rebuilt purely from the result JSON the page already holds (sent
back by the frontend together with PNG snapshots of its charts) -- nothing
is re-estimated, so exporting is instant, in the same spirit as the main
results and ML Comparison exports (pls/report.py, ml_compare/report.py).
"""

from __future__ import annotations

import base64
import io
import statistics
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from i18n import DEFAULT_LANG, t
from pls.report import _add_heading, _add_table, _fmt

POWER_THRESHOLD = 0.8
SIGNIFICANCE = 0.05
MAX_IMAGE_BYTES = 8 * 1024 * 1024


# ------------------------------------------------------------------ helpers
def _new_document(title: str, meta: str) -> Document:
    doc = Document()
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(meta)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x6B, 0x73, 0x85)
    return doc


def _note(doc: Document, text: str) -> None:
    run = doc.add_paragraph().add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x73, 0x85)


def _add_images(doc: Document, images: list | None) -> None:
    """Embeds chart snapshots sent as data: URLs; anything malformed or
    oversized is skipped rather than failing the whole export."""
    for img in images or []:
        data_url = (img or {}).get("dataUrl") or ""
        if not data_url.startswith("data:image/") or "," not in data_url:
            continue
        try:
            raw = base64.b64decode(data_url.split(",", 1)[1], validate=False)
        except (ValueError, TypeError):
            continue
        if not raw or len(raw) > MAX_IMAGE_BYTES:
            continue
        try:
            doc.add_picture(io.BytesIO(raw), width=Cm(16))
        except Exception:  # noqa: BLE001 -- not a decodable image
            continue
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(str(img.get("label") or ""))
        r.italic = True
        r.font.size = Pt(9)


def _save(doc: Document) -> io.BytesIO:
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _method_label(method: str) -> str:
    return "CB-SEM" if method == "cbsem" else "PLS-SEM"


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _quantile(sorted_vals: list[float], q: float) -> float | None:
    """Linear-interpolation quantile ("type 7"), matching the page's box plots."""
    if not sorted_vals:
        return None
    pos = (len(sorted_vals) - 1) * q
    lo = int(pos)
    rest = pos - lo
    if lo + 1 < len(sorted_vals):
        return sorted_vals[lo] + rest * (sorted_vals[lo + 1] - sorted_vals[lo])
    return sorted_vals[lo]


def _summary(values: list) -> dict | None:
    clean = sorted(float(v) for v in values if v is not None)
    if not clean:
        return None
    return {
        "mean": statistics.fmean(clean), "median": _quantile(clean, 0.5),
        "q1": _quantile(clean, 0.25), "q3": _quantile(clean, 0.75),
        "min": clean[0], "max": clean[-1], "n": len(clean),
    }


# ------------------------------------------------------------------ power
def build_power_word_report(data: dict, images: list | None = None, lang: str = DEFAULT_LANG) -> io.BytesIO:
    meta = (f"{_method_label(data.get('method'))} · n_mc = {data.get('n_mc')}"
            + (f" · bootstrap = {data['n_boot_inner']}" if data.get("n_boot_inner") else "")
            + f" · {t('rpt_export_date', lang)}: {_now()}")
    doc = _new_document(t("rpt_title_power", lang), meta)

    points = data.get("points") or []
    sample_sizes = data.get("sample_sizes") or sorted({p["n"] for p in points})
    paths: dict[str, dict] = {}
    for p in points:
        key = f"{p['source']}->{p['target']}"
        paths.setdefault(key, {"label": f"{p['source_name']} → {p['target_name']}", "points": []})["points"].append(p)

    _add_heading(doc, t("rpt_an_overview", lang), level=1)
    doc.add_paragraph(t("rpt_power_overview_text", lang, paths=len(paths), sizes=len(sample_sizes),
                        n_from=min(sample_sizes) if sample_sizes else "—",
                        n_to=max(sample_sizes) if sample_sizes else "—", n_mc=data.get("n_mc")))
    doc.add_paragraph(t("rpt_power_method_text", lang))

    _add_heading(doc, t("rpt_power_min_n_title", lang), level=1)
    rows = []
    for info in paths.values():
        pts = sorted(info["points"], key=lambda p: p["n"])
        reached = next((p["n"] for p in pts if p.get("power") is not None and p["power"] >= POWER_THRESHOLD), None)
        largest = pts[-1] if pts else {}
        rows.append([
            info["label"],
            reached if reached is not None else t("rpt_power_not_reached", lang),
            _fmt(largest.get("power")),
            _fmt(statistics.fmean([p["mean_estimate"] for p in pts if p.get("mean_estimate") is not None]))
            if any(p.get("mean_estimate") is not None for p in pts) else "—",
        ])
    _add_table(doc, [t("rpt_an_path", lang), t("rpt_power_min_n", lang),
                     t("rpt_power_at_max_n", lang, n=max(sample_sizes) if sample_sizes else "—"),
                     t("rpt_power_mean_estimate", lang)], rows)
    _note(doc, t("rpt_power_threshold_note", lang))

    if images:
        _add_heading(doc, t("rpt_an_charts", lang), level=1)
        _add_images(doc, images)

    _add_heading(doc, t("rpt_an_detail", lang), level=1)
    detail = []
    for info in paths.values():
        for p in sorted(info["points"], key=lambda p: p["n"]):
            detail.append([info["label"], p["n"], _fmt(p.get("power")),
                           f"{p.get('n_converged')}/{p.get('n_replicates')}", _fmt(p.get("mean_estimate"))])
    _add_table(doc, [t("rpt_an_path", lang), "n", "Power", t("rpt_power_converged", lang),
                     t("rpt_power_mean_estimate", lang)], detail)
    return _save(doc)


# ------------------------------------------------------------------ sensitivity
def build_sensitivity_word_report(data: dict, mode: str = "shrink", images: list | None = None,
                                  lang: str = DEFAULT_LANG) -> io.BytesIO:
    resample = mode == "resample"
    points = data.get("points") or []
    converged = [p for p in points if p.get("converged")]
    constructs = data.get("constructs") or []
    paths = data.get("paths") or []
    has_p = bool(data.get("has_p_values"))
    path_label = {p["id"]: f"{p['source_name']} → {p['target_name']}" for p in paths}

    meta = f"{_method_label(data.get('method'))} · {t('rpt_export_date', lang)}: {_now()}"
    title = t("rpt_title_sensitivity_resample" if resample else "rpt_title_sensitivity", lang)
    doc = _new_document(title, meta)

    _add_heading(doc, t("rpt_an_overview", lang), level=1)
    if resample:
        doc.add_paragraph(t("rpt_sens_resample_overview", lang, n_total=data.get("n_total"), new_n=data.get("new_n"),
                            iters=data.get("n_iterations"), conv=len(converged), total=len(points)))
    else:
        ns = [p["n"] for p in points]
        doc.add_paragraph(t("rpt_sens_shrink_overview", lang, n_total=data.get("n_total"), step=data.get("step"),
                            steps=len(points), n_min=min(ns) if ns else "—", n_max=max(ns) if ns else "—",
                            conv=len(converged), total=len(points)))
    if has_p:
        doc.add_paragraph(t("rpt_sens_pvalue_note", lang, n_boot=data.get("n_boot") or "—")
                          if data.get("method") != "cbsem" else t("rpt_sens_pvalue_note_cbsem", lang))

    if images:
        _add_heading(doc, t("rpt_an_charts", lang), level=1)
        _add_images(doc, images)

    if resample:
        stat_headers = ["Median", "Q1", "Q3", "Min", "Max", t("rpt_an_mean", lang)]

        def stat_row(label, values):
            s = _summary(values)
            return [label] + (["—"] * 6 if not s else
                              [_fmt(s["median"]), _fmt(s["q1"]), _fmt(s["q3"]), _fmt(s["min"]), _fmt(s["max"]), _fmt(s["mean"])])

        _add_heading(doc, t("rpt_sens_r2_dist", lang), level=1)
        _add_table(doc, [t("rpt_endogenous", lang)] + stat_headers,
                   [stat_row(c["name"], [p["r_squared"].get(c["id"]) for p in converged]) for c in constructs])
        _add_heading(doc, t("rpt_sens_path_dist", lang), level=1)
        _add_table(doc, [t("rpt_an_path", lang)] + stat_headers,
                   [stat_row(path_label[p["id"]], [pt["paths"].get(p["id"]) for pt in converged]) for p in paths])
        if has_p:
            _add_heading(doc, t("rpt_sens_pvalue_dist", lang), level=1)
            rows = []
            for p in paths:
                vals = [pt["p_values"].get(p["id"]) for pt in converged if pt.get("p_values")]
                vals = [v for v in vals if v is not None]
                share = f"{100 * sum(v < SIGNIFICANCE for v in vals) / len(vals):.0f}%" if vals else "—"
                rows.append(stat_row(path_label[p["id"]], vals)[:6] + [share])
            _add_table(doc, [t("rpt_an_path", lang), "Median", "Q1", "Q3", "Min", "Max",
                             t("rpt_sens_share_significant", lang)], rows)
    else:
        _add_heading(doc, t("rpt_sens_r2_by_n", lang), level=1)
        _add_table(doc, ["n", t("rpt_converged", lang)] + [c["name"] for c in constructs],
                   [[p["n"], t("rpt_yes", lang) if p.get("converged") else t("rpt_no", lang)]
                    + [_fmt(p["r_squared"].get(c["id"])) for c in constructs] for p in points])
        _add_heading(doc, t("rpt_sens_path_by_n", lang), level=1)
        for p in paths:
            _add_heading(doc, path_label[p["id"]], level=2)
            headers = ["n", t("rpt_path_coefficient", lang)] + (["p-value", t("rpt_significance", lang)] if has_p else [])
            rows = []
            for pt in points:
                coef = pt["paths"].get(p["id"])
                row = [pt["n"], _fmt(coef)]
                if has_p:
                    pv = (pt.get("p_values") or {}).get(p["id"])
                    row += [_fmt(pv, 4), "—" if pv is None else ("p < 0.05" if pv < SIGNIFICANCE else "n.s.")]
                rows.append(row)
            _add_table(doc, headers, rows)

    non_conv = len(points) - len(converged)
    if non_conv:
        _note(doc, t("rpt_sens_nonconverged", lang, n=non_conv, total=len(points)))
    return _save(doc)
