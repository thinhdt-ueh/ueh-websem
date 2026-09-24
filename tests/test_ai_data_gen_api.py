"""Tests for the /api/ai_data_gen/* routes. Never calls a real provider's
API -- monkeypatches `_call_openai`/`_call_gemini`/`_call_claude` (imported
into routes.ai_data_gen_api's own namespace from routes.ai_report_api, so
that's the module patched here, matching how the route's bare-name lookup
resolves them at call time).
"""

from __future__ import annotations

import io
import json
import os
import urllib.error

import openpyxl

import routes.ai_data_gen_api as ai_data_gen_api


def _codebook():
    return [
        {"column": "PU1", "question_text": "I find the system useful."},
        {"column": "PU2", "question_text": "Using the system improves my performance."},
        {"column": "PEOU1", "question_text": "The system is easy to use."},
    ]


def _csv_for(columns, n, value=4, age=30, gender="male", extra=None, persona="A 25-year-old tech-savvy student"):
    extra = extra or {}
    header = ",".join(["persona_description"] + list(columns) + ["resp_age", "resp_gender"] + list(extra.keys()))
    row = ",".join([f'"{persona}"'] + [str(value) for _ in columns] + [str(age), gender] + [str(v) for v in extra.values()])
    return "\n".join([header] + [row] * n)


def _demo_attrs():
    return [
        {"name": "Income", "type": "numeric", "min": 5, "max": 50},
        {"name": "Education", "type": "categorical", "options": ["High School", "Bachelor", "Master"]},
    ]


def _batch_payload(**overrides):
    payload = {
        "provider": "openai",
        "api_key": "sk-test-123",
        "model": "gpt-4o-mini",
        "temperature": 0.7,
        "system_prompt": "SYSTEM",
        "user_prompt": "USER",
        "columns": ["PU1", "PU2", "PEOU1"],
        "likert_min": 1,
        "likert_max": 5,
        "start_row": 1,
        "end_row": 5,
        "demo_age_min": 20,
        "demo_age_max": 40,
        "lang": "en",
    }
    payload.update(overrides)
    return payload


# ---------------- suggest_prompt ----------------

def test_suggest_prompt_builds_expected_shape(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(),
        "demographics": {"occupation": "university students", "target_population": "recent app users"},
        "n_rows": 45,
        "likert_scale": 5,
        "lang": "en",
    })
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert "PU1" in data["system_prompt"]
    assert "resp_age" in data["system_prompt"]
    assert "resp_gender" in data["system_prompt"]
    assert "persona_description" in data["system_prompt"]
    assert "university students" in data["system_prompt"]
    assert "recent app users" in data["system_prompt"]
    assert "45" in data["user_prompt"]
    assert data["batch_size"] == 25
    assert data["total_batches"] == 2
    assert "first_batch_instruction" in data
    assert "PU1" in data["first_batch_instruction"]
    # persona_description must be the FIRST column so the model commits to
    # a role before generating the answers that follow it in the same row.
    assert "header row: persona_description," in data["first_batch_instruction"]
    # The anti-collinearity "add ~1-point variation" instruction must be
    # paired with an explicit clamp to the declared scale bounds -- without
    # it the model can overshoot past the boundary (e.g. 5 -> 6).
    assert "MUST still stay within the valid scale range 1-5" in data["system_prompt"]


def test_suggest_prompt_respects_custom_batch_size(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(), "n_rows": 100, "likert_scale": 5, "batch_size": 25, "lang": "en",
    })
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert data["batch_size"] == 25
    assert data["total_batches"] == 4
    assert "#1-#25" in data["first_batch_instruction"]


def test_suggest_prompt_clamps_batch_size_to_allowed_range(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(), "n_rows": 100, "likert_scale": 5, "batch_size": 9999, "lang": "en",
    })
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["batch_size"] == 50


def test_suggest_prompt_rejects_empty_codebook(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={"codebook": [], "n_rows": 50, "likert_scale": 5})
    assert resp.status_code == 400


def test_suggest_prompt_rejects_duplicate_column(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": [{"column": "PU1", "question_text": "a"}, {"column": "PU1", "question_text": "b"}],
        "n_rows": 50, "likert_scale": 5,
    })
    assert resp.status_code == 400


def test_suggest_prompt_includes_qualitative_instruction_and_tag(client):
    codebook = _codebook() + [{"column": "OPEN1", "question_text": "What would you improve?", "type": "qualitative"}]
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": codebook, "n_rows": 50, "likert_scale": 5, "lang": "en",
    })
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert "open-ended" in data["system_prompt"]
    assert "OPEN1" in data["system_prompt"]
    assert "qualitative" in data["first_batch_instruction"]


def test_suggest_prompt_omits_qualitative_instruction_when_no_qualitative_columns(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(), "n_rows": 50, "likert_scale": 5, "lang": "en",
    })
    assert resp.status_code == 200, resp.get_json()
    assert "open-ended" not in resp.get_json()["system_prompt"]


# ---------------- custom demographic attributes: validation ----------------

def test_suggest_prompt_includes_custom_attrs_in_system_prompt(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(), "n_rows": 50, "likert_scale": 5, "demo_attributes": _demo_attrs(),
    })
    assert resp.status_code == 200, resp.get_json()
    prompt = resp.get_json()["system_prompt"]
    assert "Income" in prompt
    assert "Education" in prompt
    assert "Bachelor" in prompt


def test_suggest_prompt_rejects_too_many_demo_attrs(client):
    attrs = [{"name": f"Attr{i}", "type": "numeric", "min": 0, "max": 10} for i in range(7)]
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(), "n_rows": 50, "likert_scale": 5, "demo_attributes": attrs,
    })
    assert resp.status_code == 400


def test_suggest_prompt_rejects_categorical_attr_with_one_option(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(), "n_rows": 50, "likert_scale": 5,
        "demo_attributes": [{"name": "X", "type": "categorical", "options": ["only-one"]}],
    })
    assert resp.status_code == 400


def test_suggest_prompt_rejects_invalid_numeric_range(client):
    resp = client.post("/api/ai_data_gen/suggest_prompt", json={
        "codebook": _codebook(), "n_rows": 50, "likert_scale": 5,
        "demo_attributes": [{"name": "X", "type": "numeric", "min": "not-a-number", "max": 10}],
    })
    assert resp.status_code == 400


def test_validate_demo_attrs_dedupes_duplicate_names():
    """Two attributes sharing a display name must still get distinct
    columns -- collisions are resolved by _slugify_attr_name, not rejected."""
    attrs = [
        {"name": "Income", "type": "numeric", "min": 0, "max": 10},
        {"name": "Income", "type": "numeric", "min": 0, "max": 20},
    ]
    cleaned, err = ai_data_gen_api._validate_demo_attributes(attrs, "en")
    assert err is None
    assert cleaned[0]["column"] != cleaned[1]["column"]


def test_validate_codebook_preserves_construct_when_present():
    cleaned, err = ai_data_gen_api._validate_codebook(
        [{"column": "PU1", "question_text": "a", "construct": "Perceived Usefulness"}], "en",
    )
    assert err is None
    assert cleaned[0]["construct"] == "Perceived Usefulness"


def test_validate_codebook_defaults_construct_to_empty_string():
    cleaned, err = ai_data_gen_api._validate_codebook([{"column": "PU1", "question_text": "a"}], "en")
    assert err is None
    assert cleaned[0]["construct"] == ""


def test_validate_codebook_defaults_type_to_likert():
    cleaned, err = ai_data_gen_api._validate_codebook([{"column": "PU1", "question_text": "a"}], "en")
    assert err is None
    assert cleaned[0]["type"] == "likert"


def test_validate_codebook_preserves_qualitative_type():
    cleaned, err = ai_data_gen_api._validate_codebook(
        [{"column": "OPEN1", "question_text": "Why?", "type": "qualitative"}], "en",
    )
    assert err is None
    assert cleaned[0]["type"] == "qualitative"


def test_validate_codebook_rejects_unknown_type_by_falling_back_to_likert():
    cleaned, err = ai_data_gen_api._validate_codebook(
        [{"column": "PU1", "question_text": "a", "type": "not-a-real-type"}], "en",
    )
    assert err is None
    assert cleaned[0]["type"] == "likert"


def test_split_codebook_columns_separates_by_type_and_preserves_order():
    codebook = [
        {"column": "PU1", "question_text": "a", "construct": "", "type": "likert"},
        {"column": "OPEN1", "question_text": "b", "construct": "", "type": "qualitative"},
        {"column": "PU2", "question_text": "c", "construct": "", "type": "likert"},
    ]
    likert_columns, qual_columns = ai_data_gen_api._split_codebook_columns(codebook)
    assert likert_columns == ["PU1", "PU2"]
    assert qual_columns == ["OPEN1"]


# ---------------- suggest_constructs ----------------

def _construct_search_json(constructs):
    return json.dumps({
        "constructs": [
            {
                "name": name,
                "theory": {"citation_apa": f"Author, A. ({name}). A study. Journal.", "doi": ""},
                "items": [{"column": col, "question_text": q} for col, q in items],
            }
            for name, items in constructs
        ],
    })


def _construct_search_payload(**overrides):
    payload = {
        "provider": "openai",
        "api_key": "sk-test-123",
        "model": "gpt-4o-mini",
        "temperature": 0.7,
        "topic": "Students' intention to use mobile wallets, based on an extended TAM",
        "n_constructs": 3,
        "lang": "en",
    }
    payload.update(overrides)
    return payload


def test_suggest_constructs_success(client, monkeypatch):
    body = _construct_search_json([
        ("Perceived Usefulness", [("PU1", "Using it improves my performance."), ("PU2", "It is useful.")]),
        ("Perceived Ease of Use", [("PEOU1", "It is easy to use.")]),
    ])
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: body)
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload())
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    constructs = data["constructs"]
    assert len(constructs) == 2
    assert constructs[0]["name"] == "Perceived Usefulness"
    assert constructs[0]["items"] == [
        {"column": "PU1", "question_text": "Using it improves my performance."},
        {"column": "PU2", "question_text": "It is useful."},
    ]
    assert constructs[0]["theory"]["citation_apa"]
    assert constructs[0]["theory"]["doi"] == ""


def test_suggest_constructs_rejects_missing_theory_citation(client, monkeypatch):
    body = json.dumps({"constructs": [{"name": "A", "items": [{"column": "A1", "question_text": "q"}]}]})
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: body)
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload())
    assert resp.status_code == 422


def test_suggest_constructs_wraps_markdown_code_fence(client, monkeypatch):
    body = _construct_search_json([("PU", [("PU1", "q1")])])
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: f"```json\n{body}\n```")
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload())
    assert resp.status_code == 200, resp.get_json()


def test_suggest_constructs_dedupes_duplicate_columns_across_constructs(client, monkeypatch):
    body = _construct_search_json([
        ("A", [("X1", "question a")]),
        ("B", [("X1", "question b")]),
    ])
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: body)
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload())
    assert resp.status_code == 200, resp.get_json()
    constructs = resp.get_json()["constructs"]
    columns = [item["column"] for group in constructs for item in group["items"]]
    assert len(columns) == len(set(columns))
    assert "X1" in columns and "X1_2" in columns


def test_suggest_constructs_malformed_json_rejected(client, monkeypatch):
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: "not json at all")
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload())
    assert resp.status_code == 422


def test_suggest_constructs_missing_topic_rejected(client):
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload(topic=""))
    assert resp.status_code == 400


def test_suggest_constructs_missing_api_key_rejected(client):
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload(api_key=""))
    assert resp.status_code == 400


def test_suggest_constructs_provider_http_error_maps(client, monkeypatch):
    def raise_401(*a, **k):
        raise urllib.error.HTTPError("url", 401, "unauthorized", {}, io.BytesIO(b"{}"))

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", raise_401)
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload())
    assert resp.status_code == 400


def test_suggest_constructs_clamps_n_constructs(client, monkeypatch):
    captured = {}

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        captured["system_msg"] = system_msg
        return _construct_search_json([("A", [("A1", "q")])])

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/suggest_constructs", json=_construct_search_payload(n_constructs=999))
    assert resp.status_code == 200, resp.get_json()
    assert str(ai_data_gen_api.MAX_CONSTRUCTS) in captured["system_msg"]


# ---------------- suggest_paths ----------------

def _tam_constructs():
    return [
        {"id": "peou", "name": "Perceived Ease of Use", "mode": "A", "indicators": ["PEOU1", "PEOU2"]},
        {"id": "pu", "name": "Perceived Usefulness", "mode": "A", "indicators": ["PU1", "PU2"]},
        {"id": "att", "name": "Attitude", "mode": "A", "indicators": ["ATT1", "ATT2"]},
        {"id": "int", "name": "Behavioral Intention", "mode": "A", "indicators": ["INT1", "INT2"]},
    ]


def _paths_search_payload(**overrides):
    payload = {
        "provider": "openai",
        "api_key": "sk-test-123",
        "model": "gpt-4o-mini",
        "temperature": 0.7,
        "constructs": _tam_constructs(),
        "system_prompt": "You are a structural-model expert.",
        "user_prompt": "Please propose a structural model for the constructs above.",
        "lang": "en",
    }
    payload.update(overrides)
    return payload


def _paths_search_json(paths, rationale="Because established TAM theory says so.", moderator_suggestions=None):
    body = {"paths": paths, "rationale": rationale}
    if moderator_suggestions is not None:
        body["moderator_suggestions"] = moderator_suggestions
    return json.dumps(body)


def test_suggest_paths_success(client, monkeypatch):
    tam_paths = [
        {"source": "peou", "target": "pu"},
        {"source": "peou", "target": "att"},
        {"source": "pu", "target": "att"},
        {"source": "pu", "target": "int"},
        {"source": "att", "target": "int"},
    ]
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: _paths_search_json(tam_paths))
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert data["paths"] == tam_paths
    assert data["rationale"]
    assert data["moderator_suggestions"] == []


def test_suggest_paths_returns_moderator_suggestions(client, monkeypatch):
    tam_paths = [{"source": "peou", "target": "pu"}]
    suggestions = [{"construct_id": "att", "reason": "Attitude may moderate PU->Intention instead of mediating it."}]
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _paths_search_json(tam_paths, moderator_suggestions=suggestions),
    )
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["moderator_suggestions"] == suggestions


def test_suggest_paths_drops_moderator_suggestion_for_unknown_construct(client, monkeypatch):
    tam_paths = [{"source": "peou", "target": "pu"}]
    suggestions = [{"construct_id": "does-not-exist", "reason": "hallucinated id"}]
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _paths_search_json(tam_paths, moderator_suggestions=suggestions),
    )
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["moderator_suggestions"] == []


def test_suggest_paths_drops_moderator_suggestion_for_interaction_construct(client, monkeypatch):
    constructs = _tam_constructs() + [
        {"id": "mod", "name": "PEOU x PU", "mode": "I", "interaction_of": ["peou", "pu"]},
    ]
    tam_paths = [{"source": "peou", "target": "pu"}, {"source": "mod", "target": "att"},
                 {"source": "peou", "target": "att"}, {"source": "pu", "target": "att"}]
    suggestions = [{"construct_id": "mod", "reason": "already an interaction construct"}]
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _paths_search_json(tam_paths, moderator_suggestions=suggestions),
    )
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload(constructs=constructs))
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["moderator_suggestions"] == []


def test_suggest_paths_wraps_markdown_code_fence(client, monkeypatch):
    body = _paths_search_json([{"source": "peou", "target": "pu"}])
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: f"```json\n{body}\n```")
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 200, resp.get_json()


def test_suggest_paths_resolves_construct_name_used_instead_of_id(client, monkeypatch):
    # A model sometimes echoes a construct's NAME instead of its id in
    # source/target despite the prompt insisting on the id -- this used to
    # hard-fail as "path references a construct that doesn't exist" even
    # though the intended model was perfectly valid.
    paths_by_name = [
        {"source": "Perceived Ease of Use", "target": "pu"},
        {"source": "pu", "target": "Attitude"},
    ]
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: _paths_search_json(paths_by_name))
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["paths"] == [{"source": "peou", "target": "pu"}, {"source": "pu", "target": "att"}]


def test_suggest_paths_resolves_construct_name_in_moderator_suggestion(client, monkeypatch):
    tam_paths = [{"source": "peou", "target": "pu"}]
    suggestions = [{"construct_id": "Attitude", "reason": "May moderate instead of mediate."}]
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _paths_search_json(tam_paths, moderator_suggestions=suggestions),
    )
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["moderator_suggestions"] == [{"construct_id": "att", "reason": "May moderate instead of mediate."}]


def test_suggest_paths_rejects_cycle(client, monkeypatch):
    cyclic = [{"source": "peou", "target": "pu"}, {"source": "pu", "target": "peou"}]
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: _paths_search_json(cyclic))
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 422
    assert "cycle" in resp.get_json()["error"].lower()


def test_suggest_paths_rejects_interaction_as_target(client, monkeypatch):
    constructs = _tam_constructs() + [
        {"id": "mod", "name": "PEOU x PU", "mode": "I", "interaction_of": ["peou", "pu"]},
    ]
    # "mod" (an interaction/moderation construct) can never receive an
    # incoming path -- proves Model.from_json's real rule is enforced here,
    # not a hand-rolled approximation of it.
    bad_paths = [{"source": "peou", "target": "pu"}, {"source": "pu", "target": "mod"}]
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: _paths_search_json(bad_paths))
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload(constructs=constructs))
    assert resp.status_code == 422
    assert "incoming path" in resp.get_json()["error"].lower()


def test_suggest_paths_missing_rationale_rejected(client, monkeypatch):
    body = json.dumps({"paths": [{"source": "peou", "target": "pu"}], "rationale": ""})
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: body)
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 422


def test_suggest_paths_missing_api_key_rejected(client):
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload(api_key=""))
    assert resp.status_code == 400


def test_suggest_paths_too_few_constructs_rejected(client):
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload(constructs=_tam_constructs()[:1]))
    assert resp.status_code == 400


def test_suggest_paths_provider_http_error_maps(client, monkeypatch):
    def raise_401(*a, **k):
        raise urllib.error.HTTPError("url", 401, "unauthorized", {}, io.BytesIO(b"{}"))

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", raise_401)
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload())
    assert resp.status_code == 400


def test_suggest_paths_missing_system_prompt_rejected(client):
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload(system_prompt=""))
    assert resp.status_code == 400


def test_suggest_paths_missing_user_prompt_rejected(client):
    resp = client.post("/api/ai_data_gen/suggest_paths", json=_paths_search_payload(user_prompt=""))
    assert resp.status_code == 400


# ---------------- suggest_paths_prompt ----------------

def test_suggest_paths_prompt_returns_prompt_text_with_no_ai_call(client, monkeypatch):
    def fail_if_called(*a, **k):
        raise AssertionError("suggest_paths_prompt must not call the AI provider")

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fail_if_called)
    monkeypatch.setattr(ai_data_gen_api, "_call_gemini", fail_if_called)
    monkeypatch.setattr(ai_data_gen_api, "_call_claude", fail_if_called)
    resp = client.post(
        "/api/ai_data_gen/suggest_paths_prompt",
        json={"constructs": _tam_constructs(), "extra_context": "B2B SaaS adoption study", "lang": "en"},
    )
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert data["system_prompt"]
    assert data["user_prompt"]
    assert "B2B SaaS adoption study" in data["user_prompt"]


def test_suggest_paths_prompt_too_few_constructs_rejected(client):
    resp = client.post(
        "/api/ai_data_gen/suggest_paths_prompt",
        json={"constructs": _tam_constructs()[:1], "lang": "en"},
    )
    assert resp.status_code == 400


def test_suggest_paths_prompt_includes_indicator_descriptions(client):
    resp = client.post(
        "/api/ai_data_gen/suggest_paths_prompt",
        json={
            "constructs": _tam_constructs(),
            "indicator_descriptions": {"PEOU1": "Learning to use the system would be easy for me"},
            "lang": "en",
        },
    )
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert "Learning to use the system would be easy for me" in data["system_prompt"]


def test_suggest_paths_prompt_mentions_moderator_possibility(client):
    resp = client.post(
        "/api/ai_data_gen/suggest_paths_prompt",
        json={"constructs": _tam_constructs(), "lang": "en"},
    )
    assert resp.status_code == 200, resp.get_json()
    assert "moderat" in resp.get_json()["system_prompt"].lower()


# ---------------- batch ----------------

def test_batch_success_single_attempt(client, monkeypatch):
    calls = []

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        calls.append(1)
        return _csv_for(["PU1", "PU2", "PEOU1"], 5)

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    rows = data["rows"]
    assert len(rows) == 5
    assert all(set(r.keys()) == {"persona_description", "PU1", "PU2", "PEOU1", "resp_age", "resp_gender"} for r in rows)
    assert all(r["PU1"] == 4 for r in rows)
    assert all(r["resp_age"] == 30 for r in rows)
    assert all(r["resp_gender"] == "male" for r in rows)
    assert all(r["persona_description"] for r in rows)
    assert len(calls) == 1
    assert "used_system_prompt" in data and "used_user_prompt" in data


def test_batch_retries_on_malformed_csv_then_succeeds(client, monkeypatch):
    calls = []

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        calls.append(1)
        if len(calls) == 1:
            return "this is not a csv at all, sorry!"
        return _csv_for(["PU1", "PU2", "PEOU1"], 5)

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 200, resp.get_json()
    assert len(resp.get_json()["rows"]) == 5
    assert len(calls) == 2


def test_batch_retry_exhaustion_returns_422(client, monkeypatch):
    calls = []

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        calls.append(1)
        return "still not a csv"

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422
    assert len(calls) == ai_data_gen_api.MAX_BATCH_ATTEMPTS


def test_batch_over_count_csv_is_truncated_and_accepted(client, monkeypatch):
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _csv_for(["PU1", "PU2", "PEOU1"], 7),
    )
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 200, resp.get_json()
    assert len(resp.get_json()["rows"]) == 5


def test_batch_under_count_csv_triggers_retry_then_fails(client, monkeypatch):
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _csv_for(["PU1", "PU2", "PEOU1"], 2),
    )
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422


def test_batch_small_shortfall_accepted_as_partial_batch(client, monkeypatch):
    """A near-miss (e.g. 24/25) is the model slightly under-counting, not a
    real error -- accepted outright rather than discarded and retried, so
    the caller's adaptive loop can just ask for the few remaining rows."""
    calls = []

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        calls.append(1)
        return _csv_for(["PU1", "PU2", "PEOU1"], 4)  # 4 of 5 requested

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(start_row=1, end_row=5))
    assert resp.status_code == 200, resp.get_json()
    assert len(resp.get_json()["rows"]) == 4
    assert len(calls) == 1  # accepted on the first attempt, no retry needed


def test_batch_row_shortfall_gets_actionable_corrective_note(client, monkeypatch):
    """A response cut short (fewer rows than asked) is almost always an
    output-length truncation, not a formatting mistake -- the corrective
    retry should say so and tell the model to shrink persona_description,
    not just repeat the generic 'invalid, try again' note."""
    captured_system_msgs = []

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        captured_system_msgs.append(system_msg)
        return _csv_for(["PU1", "PU2", "PEOU1"], 2)  # always short of expected_n=5

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422
    assert len(captured_system_msgs) == ai_data_gen_api.MAX_BATCH_ATTEMPTS
    assert "CUT OFF" in captured_system_msgs[1]
    assert "6 words max" in captured_system_msgs[1]


def test_batch_out_of_range_value_rejected(client, monkeypatch):
    header = "persona_description,PU1,PU2,PEOU1,resp_age,resp_gender"
    bad_row = '"A persona",9,9,9,30,male'
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: "\n".join([header] + [bad_row] * 5),
    )
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422


def test_batch_range_violation_gets_actionable_corrective_note(client, monkeypatch):
    """An out-of-range value is almost always the anti-collinearity 'add
    ~1-point variation' instruction overshooting the boundary -- the
    corrective retry should name that mechanism and tell the model to clamp
    at the edge, not just repeat the generic 'invalid, try again' note."""
    header = "persona_description,PU1,PU2,PEOU1,resp_age,resp_gender"
    bad_row = '"A persona",9,4,4,30,male'
    captured_system_msgs = []

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        captured_system_msgs.append(system_msg)
        return "\n".join([header] + [bad_row] * 5)

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422
    assert len(captured_system_msgs) == ai_data_gen_api.MAX_BATCH_ATTEMPTS
    assert "OUTSIDE the 1-5 range" in captured_system_msgs[1]
    assert "KEEP that edge value" in captured_system_msgs[1]


def test_batch_empty_persona_description_rejected(client, monkeypatch):
    header = "persona_description,PU1,PU2,PEOU1,resp_age,resp_gender"
    bad_row = ',4,4,4,30,male'
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: "\n".join([header] + [bad_row] * 5),
    )
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422


def test_batch_out_of_range_age_rejected(client, monkeypatch):
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _csv_for(["PU1", "PU2", "PEOU1"], 5, age=99),  # outside demo_age_min/max 20-40
    )
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422


def test_batch_invalid_gender_rejected(client, monkeypatch):
    monkeypatch.setattr(
        ai_data_gen_api, "_call_openai",
        lambda *a, **k: _csv_for(["PU1", "PU2", "PEOU1"], 5, gender="unknown"),
    )
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 422


def test_batch_provider_http_error_maps_and_does_not_retry(client, monkeypatch):
    calls = []

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        calls.append(1)
        raise urllib.error.HTTPError("url", 401, "unauthorized", hdrs=None, fp=None)

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 400
    assert "invalid" in resp.get_json()["error"].lower() or "hết hạn" in resp.get_json()["error"]
    assert len(calls) == 1


def test_batch_missing_api_key_rejected(client):
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(api_key=""))
    assert resp.status_code == 400


def test_batch_wraps_markdown_code_fence(client, monkeypatch):
    csv_text = _csv_for(["PU1", "PU2", "PEOU1"], 5)
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: f"```csv\n{csv_text}\n```")
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload())
    assert resp.status_code == 200, resp.get_json()
    assert len(resp.get_json()["rows"]) == 5


def test_batch_accepts_up_to_50_rows(client, monkeypatch):
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: _csv_for(["PU1", "PU2", "PEOU1"], 50))
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(start_row=1, end_row=50))
    assert resp.status_code == 200, resp.get_json()
    assert len(resp.get_json()["rows"]) == 50


def test_batch_rejects_more_than_50_rows(client):
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(start_row=1, end_row=51))
    assert resp.status_code == 400


# ---------------- qualitative columns: batch ----------------

def test_batch_with_qualitative_column_success(client, monkeypatch):
    header = "persona_description,PU1,PU2,OPEN1,resp_age,resp_gender"
    row = '"A 25-year-old student",4,4,"I find the app fast and reliable.",30,male'
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: "\n".join([header] + [row] * 5))
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(
        columns=["PU1", "PU2", "OPEN1"], qualitative_columns=["OPEN1"],
    ))
    assert resp.status_code == 200, resp.get_json()
    rows = resp.get_json()["rows"]
    assert len(rows) == 5
    assert all(r["OPEN1"] == "I find the app fast and reliable." for r in rows)
    assert all(r["PU1"] == 4 for r in rows)
    assert all(set(r.keys()) == {"persona_description", "PU1", "PU2", "OPEN1", "resp_age", "resp_gender"} for r in rows)


def test_batch_empty_qualitative_answer_rejected(client, monkeypatch):
    header = "persona_description,PU1,OPEN1,resp_age,resp_gender"
    row = '"A persona",4,"",30,male'
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: "\n".join([header] + [row] * 5))
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(
        columns=["PU1", "OPEN1"], qualitative_columns=["OPEN1"],
    ))
    assert resp.status_code == 422


def test_batch_qualitative_column_ignored_when_not_declared(client, monkeypatch):
    """With no qualitative_columns sent, every column is validated as a
    Likert integer -- free text in that column must be rejected, matching
    pre-existing behavior for a codebook with no qualitative rows at all."""
    header = "persona_description,PU1,OPEN1,resp_age,resp_gender"
    row = '"A persona",4,"free text answer",30,male'
    monkeypatch.setattr(ai_data_gen_api, "_call_openai", lambda *a, **k: "\n".join([header] + [row] * 5))
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(columns=["PU1", "OPEN1"]))
    assert resp.status_code == 422


# ---------------- custom demographic attributes: batch ----------------

def test_batch_with_custom_attributes_success(client, monkeypatch):
    cleaned, err = ai_data_gen_api._validate_demo_attributes(_demo_attrs(), "en")
    assert err is None
    income_col, edu_col = cleaned[0]["column"], cleaned[1]["column"]

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        return _csv_for(["PU1", "PU2", "PEOU1"], 5, extra={income_col: 20, edu_col: "bachelor"})

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(demo_attributes=_demo_attrs()))
    assert resp.status_code == 200, resp.get_json()
    rows = resp.get_json()["rows"]
    assert len(rows) == 5
    assert all(r[income_col] == 20 for r in rows)
    # normalized to the declared canonical spelling ("Bachelor"), not the
    # AI's lowercase "bachelor"
    assert all(r[edu_col] == "Bachelor" for r in rows)


def test_batch_custom_numeric_out_of_range_rejected(client, monkeypatch):
    cleaned, _ = ai_data_gen_api._validate_demo_attributes(_demo_attrs(), "en")
    income_col, edu_col = cleaned[0]["column"], cleaned[1]["column"]

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        return _csv_for(["PU1", "PU2", "PEOU1"], 5, extra={income_col: 999, edu_col: "Bachelor"})

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(demo_attributes=_demo_attrs()))
    assert resp.status_code == 422


def test_batch_custom_categorical_invalid_value_rejected(client, monkeypatch):
    cleaned, _ = ai_data_gen_api._validate_demo_attributes(_demo_attrs(), "en")
    income_col, edu_col = cleaned[0]["column"], cleaned[1]["column"]

    def fake_call(api_key, model, system_msg, user_msg, temperature):
        return _csv_for(["PU1", "PU2", "PEOU1"], 5, extra={income_col: 20, edu_col: "PhD"})

    monkeypatch.setattr(ai_data_gen_api, "_call_openai", fake_call)
    resp = client.post("/api/ai_data_gen/batch", json=_batch_payload(demo_attributes=_demo_attrs()))
    assert resp.status_code == 422


# ---------------- finalize ----------------

def _rows(n, columns=("PU1", "PU2", "PEOU1"), value=4, age=30, gender="male"):
    out = []
    for i in range(n):
        row = {"persona_description": f"Respondent {i}: a {age}-year-old with a consistent attitude"}
        row.update({c: value for c in columns})
        row["resp_age"] = age
        row["resp_gender"] = gender if i % 2 == 0 else ("female" if gender == "male" else "male")
        out.append(row)
    return out


def _varied_rows(n, columns):
    """Unlike _rows() (a constant value, fine for shape/count assertions),
    real PLS-SEM estimation rejects zero-variance indicators -- this cycles
    through 1-5 per row so /api/analyze can actually run on the data."""
    out = []
    for i in range(n):
        row = {"persona_description": f"Respondent {i}: a consistent persona"}
        row.update({c: 1 + (i + j) % 5 for j, c in enumerate(columns)})
        row["resp_age"] = 20 + (i % 30)
        row["resp_gender"] = "male" if i % 2 == 0 else "female"
        out.append(row)
    return out


def _finalize_payload(columns, rows, **overrides):
    payload = {
        "filename": "synthetic.csv",
        "columns": columns,
        "rows": rows,
        "codebook": [{"column": c, "question_text": f"Question about {c}"} for c in columns],
        "demographics": {
            "occupation": "university students",
            "location": "Ho Chi Minh City",
            "target_population": "Students who use mobile wallets",
            "age_min": 20, "age_max": 40, "gender_mix": "balanced",
        },
        "provider": "openai",
        "model": "gpt-4o-mini",
        "temperature": 0.7,
        "likert_scale": 5,
        "batches": [
            {"start_row": 1, "end_row": len(rows), "system_prompt": "SYS", "user_prompt": "USER"},
        ],
        "lang": "en",
    }
    payload.update(overrides)
    return payload


def test_finalize_success_matches_upload_shape_plus_new_keys(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1", "PU2", "PEOU1"], _rows(40)))
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert set(data.keys()) == {
        "file_id", "filename", "columns", "numeric_columns", "n_rows", "preview",
        "demographics_summary", "demo_attributes", "descriptive_stats",
    }
    assert data["n_rows"] == 40
    assert data["columns"] == ["PU1", "PU2", "PEOU1"]
    assert len(data["preview"]) == 10
    # indicator/demographic columns must never mix
    assert "resp_age" not in data["columns"]
    assert "resp_gender" not in data["columns"]


def test_finalize_descriptive_stats_are_correct(client):
    rows = _rows(30, columns=("PU1",), value=4, age=25)
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], rows))
    assert resp.status_code == 200, resp.get_json()
    stats = resp.get_json()["descriptive_stats"]
    assert stats["indicators"]["PU1"]["mean"] == 4.0
    assert stats["demographics"]["age"]["mean"] == 25.0
    assert stats["demographics"]["gender"]["male"]["count"] == 15
    assert stats["demographics"]["gender"]["female"]["count"] == 15
    assert stats["demographics"]["gender"]["male"]["pct"] == 50.0


def test_finalize_echoes_demographics_summary(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], _rows(30, columns=("PU1",))))
    summary = resp.get_json()["demographics_summary"]
    assert summary["occupation"] == "university students"
    assert summary["target_population"] == "Students who use mobile wallets"


def test_finalize_persists_ai_meta_files(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], _rows(30, columns=("PU1",))))
    file_id = resp.get_json()["file_id"]
    with client.application.test_request_context():
        meta, demo_df = ai_data_gen_api._load_ai_gen_metadata(file_id)
    assert meta is not None
    assert meta["demographics"]["occupation"] == "university students"
    assert meta["codebook"][0]["column"] == "PU1"
    assert meta["batches"][0]["system_prompt"] == "SYS"
    assert len(demo_df) == 30
    assert set(demo_df.columns) == {"respondent_id", "persona_description", "resp_age", "resp_gender"}
    assert demo_df["respondent_id"].tolist() == [f"R{i + 1:04d}" for i in range(30)]


def test_finalize_rejects_too_few_rows(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], _rows(5, columns=("PU1",))))
    assert resp.status_code == 400


def test_finalize_rejects_too_many_rows(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], _rows(501, columns=("PU1",))))
    assert resp.status_code == 400


def test_finalize_rejects_shape_mismatch(client):
    rows = _rows(40)
    rows[0] = {"PU1": 4, "PU2": 4, "resp_age": 30, "resp_gender": "male"}  # missing PEOU1
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1", "PU2", "PEOU1"], rows))
    assert resp.status_code == 400


def test_finalize_rejects_invalid_gender_value(client):
    rows = _rows(40)
    rows[0]["resp_gender"] = "unknown"
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1", "PU2", "PEOU1"], rows))
    assert resp.status_code == 400


def test_finalize_result_resolvable_by_analyze(client, tam_model_json):
    """The critical proof that downstream code needs zero changes: a
    finalized AI-generated dataset's file_id must resolve through
    /api/analyze's existing prefix-match file lookup exactly like an
    uploaded or sample-loaded file does -- and demographic columns must
    never leak into the model's usable columns."""
    columns = ["PEOU1", "PEOU2", "PEOU3", "PU1", "PU2", "PU3", "ATT1", "ATT2", "ATT3", "INT1", "INT2", "INT3"]
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(columns, _varied_rows(60, columns)))
    assert resp.status_code == 200, resp.get_json()
    file_id = resp.get_json()["file_id"]

    analyze_resp = client.post("/api/analyze", json={"file_id": file_id, "model": tam_model_json, "lang": "en"})
    assert analyze_resp.status_code == 200, analyze_resp.get_json()


# ---------------- custom demographic attributes: finalize ----------------

def test_finalize_persists_and_computes_custom_attrs(client):
    cleaned, err = ai_data_gen_api._validate_demo_attributes(_demo_attrs(), "en")
    assert err is None
    income_col, edu_col = cleaned[0]["column"], cleaned[1]["column"]

    rows = _rows(30, columns=("PU1",))
    for i, row in enumerate(rows):
        row[income_col] = 10 + (i % 5)
        row[edu_col] = "bachelor" if i % 2 == 0 else "Master"

    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        ["PU1"], rows, demo_attributes=_demo_attrs(),
    ))
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()

    assert data["demo_attributes"][0]["column"] == income_col
    assert data["demo_attributes"][1]["options"] == ["High School", "Bachelor", "Master"]

    custom_stats = data["descriptive_stats"]["demographics"]["custom"]
    income_stats = next(c for c in custom_stats if c["column"] == income_col)
    edu_stats = next(c for c in custom_stats if c["column"] == edu_col)
    assert income_stats["type"] == "numeric"
    assert edu_stats["type"] == "categorical"
    # normalized to canonical spelling ("Bachelor"), counted even though the
    # row data used lowercase "bachelor"
    assert edu_stats["counts"]["Bachelor"]["count"] == 15
    assert edu_stats["counts"]["Master"]["count"] == 15
    assert edu_stats["counts"]["High School"]["count"] == 0

    file_id = data["file_id"]
    with client.application.test_request_context():
        meta, demo_df = ai_data_gen_api._load_ai_gen_metadata(file_id)
    assert meta["demo_attributes"] == cleaned
    assert income_col in demo_df.columns
    assert edu_col in demo_df.columns
    assert demo_df[edu_col].tolist()[0] == "Bachelor"

    # Custom attribute columns must never leak into the SEM indicator file,
    # exactly like age/gender.
    dl = client.get(f"/api/ai_data_gen/download?file_id={file_id}")
    assert income_col.encode() not in dl.data
    assert edu_col.encode() not in dl.data


def test_finalize_rejects_invalid_custom_categorical_value(client):
    rows = _rows(30, columns=("PU1",))
    cleaned, _ = ai_data_gen_api._validate_demo_attributes(_demo_attrs(), "en")
    income_col, edu_col = cleaned[0]["column"], cleaned[1]["column"]
    for row in rows:
        row[income_col] = 20
        row[edu_col] = "PhD"  # not one of the declared options
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        ["PU1"], rows, demo_attributes=_demo_attrs(),
    ))
    assert resp.status_code == 400


# ---------------- download ----------------

def test_download_resolves_by_file_id_after_finalize(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], _rows(40, columns=("PU1",))))
    file_id = resp.get_json()["file_id"]
    dl = client.get(f"/api/ai_data_gen/download?file_id={file_id}")
    assert dl.status_code == 200
    assert b"PU1" in dl.data
    assert b"resp_age" not in dl.data
    assert b"persona_description" not in dl.data


def test_finalize_persists_persona_description_in_respondents_file_only(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], _rows(30, columns=("PU1",))))
    assert resp.status_code == 200, resp.get_json()
    file_id = resp.get_json()["file_id"]
    with client.application.test_request_context():
        _meta, demo_df = ai_data_gen_api._load_ai_gen_metadata(file_id)
    assert "persona_description" in demo_df.columns
    assert all(demo_df["persona_description"].str.len() > 0)


def test_download_missing_file_id_404(client):
    resp = client.get("/api/ai_data_gen/download?file_id=doesnotexist")
    assert resp.status_code == 404


# ---------------- qualitative columns: finalize ----------------

def _rows_with_qual(n, likert_columns=("PU1", "PU2"), qual_column="OPEN1", value=4, age=30, gender="male"):
    out = []
    for i in range(n):
        row = {"persona_description": f"Respondent {i}: a {age}-year-old with a consistent attitude"}
        row.update({c: value for c in likert_columns})
        row[qual_column] = f"Free-text answer from respondent {i}."
        row["resp_age"] = age
        row["resp_gender"] = gender if i % 2 == 0 else ("female" if gender == "male" else "male")
        out.append(row)
    return out


def _qual_codebook(likert_columns=("PU1", "PU2"), qual_column="OPEN1"):
    codebook = [{"column": c, "question_text": f"Question about {c}"} for c in likert_columns]
    codebook.append({"column": qual_column, "question_text": "What would you improve?", "type": "qualitative"})
    return codebook


def test_finalize_keeps_qualitative_column_out_of_indicator_dataset(client):
    columns = ["PU1", "PU2", "OPEN1"]
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        columns, _rows_with_qual(30), codebook=_qual_codebook(),
    ))
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert data["columns"] == ["PU1", "PU2"]
    assert data["numeric_columns"] == ["PU1", "PU2"]
    assert "OPEN1" not in data["columns"]
    assert all("OPEN1" not in row for row in data["preview"])


def test_finalize_persists_qualitative_answers_in_respondents_file(client):
    columns = ["PU1", "PU2", "OPEN1"]
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        columns, _rows_with_qual(30), codebook=_qual_codebook(),
    ))
    assert resp.status_code == 200, resp.get_json()
    file_id = resp.get_json()["file_id"]
    with client.application.test_request_context():
        _meta, demo_df = ai_data_gen_api._load_ai_gen_metadata(file_id)
    assert "OPEN1" in demo_df.columns
    assert all(demo_df["OPEN1"].str.startswith("Free-text answer"))


def test_finalize_rejects_empty_qualitative_answer(client):
    rows = _rows_with_qual(30)
    rows[0]["OPEN1"] = "   "
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        ["PU1", "PU2", "OPEN1"], rows, codebook=_qual_codebook(),
    ))
    assert resp.status_code == 400


def test_export_includes_qualitative_column_in_respondent_profile_sheet(client):
    columns = ["PU1", "PU2", "OPEN1"]
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        columns, _rows_with_qual(30), codebook=_qual_codebook(),
    ))
    file_id = resp.get_json()["file_id"]
    exp = client.get(f"/api/ai_data_gen/export?file_id={file_id}")
    assert exp.status_code == 200, exp.get_data()
    wb = openpyxl.load_workbook(io.BytesIO(exp.data))
    survey_ws = wb["Survey Data"]
    survey_header = [c.value for c in next(survey_ws.iter_rows(min_row=1, max_row=1))]
    assert "OPEN1" not in survey_header

    profile_ws = wb["Respondent Profile"]
    profile_values = [[c.value for c in row] for row in profile_ws.iter_rows()]
    flat_text = [str(v) for row in profile_values for v in row if v is not None]
    assert any(v == "OPEN1" for v in flat_text)
    assert any(str(v).startswith("Free-text answer") for v in flat_text)


# ---------------- export ----------------

def test_export_404_for_non_ai_generated_file(client):
    csv_bytes = b"PU1,PU2\n1,2\n3,4\n"
    up = client.post(
        "/api/upload",
        data={"file": (io.BytesIO(csv_bytes), "plain.csv"), "lang": "en"},
        content_type="multipart/form-data",
    )
    file_id = up.get_json()["file_id"]
    resp = client.get(f"/api/ai_data_gen/export?file_id={file_id}")
    assert resp.status_code == 404


def test_export_404_unknown_file_id(client):
    resp = client.get("/api/ai_data_gen/export?file_id=doesnotexist")
    assert resp.status_code == 404


def test_export_returns_four_sheet_workbook(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1", "PU2"], _rows(30, columns=("PU1", "PU2"))))
    file_id = resp.get_json()["file_id"]

    export_resp = client.get(f"/api/ai_data_gen/export?file_id={file_id}")
    assert export_resp.status_code == 200
    assert export_resp.content_type.startswith("application/vnd.openxmlformats")

    wb = openpyxl.load_workbook(io.BytesIO(export_resp.data))
    assert wb.sheetnames == ["Survey Data", "Respondent Profile", "Descriptive Statistics", "Prompt Transparency"]

    data_ws = wb["Survey Data"]
    assert [c.value for c in next(data_ws.iter_rows(min_row=1, max_row=1))] == ["respondent_id", "PU1", "PU2"]
    assert data_ws.max_row == 31  # header + 30 rows
    survey_ids = [row[0].value for row in data_ws.iter_rows(min_row=2, max_row=31)]
    assert survey_ids == [f"R{i + 1:04d}" for i in range(30)]

    profile_ws = wb["Respondent Profile"]
    profile_values = [row[1].value for row in profile_ws.iter_rows(min_row=1, max_row=3)]
    assert "university students" in profile_values

    # The join key must line up identically between the two sheets, in the
    # same row order, so the two exported artifacts can be cross-referenced.
    profile_header_row = next(
        r for r in profile_ws.iter_rows(min_row=1, max_row=profile_ws.max_row)
        if r[0].value == "respondent_id"
    )
    header_row_idx = profile_header_row[0].row
    profile_ids = [
        row[0].value
        for row in profile_ws.iter_rows(min_row=header_row_idx + 1, max_row=profile_ws.max_row)
    ]
    assert profile_ids == survey_ids

    transparency_ws = wb["Prompt Transparency"]
    header_row = [c.value for c in next(transparency_ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == ["Rows", "Provider", "Model", "Temperature", "System Prompt", "User Prompt"]
    assert transparency_ws.cell(row=2, column=5).value == "SYS"


def test_export_includes_custom_demographic_attributes(client):
    cleaned, _ = ai_data_gen_api._validate_demo_attributes(_demo_attrs(), "en")
    income_col, edu_col = cleaned[0]["column"], cleaned[1]["column"]
    rows = _rows(30, columns=("PU1", "PU2"))
    for i, row in enumerate(rows):
        row[income_col] = 10 + (i % 5)
        row[edu_col] = "Bachelor" if i % 2 == 0 else "Master"

    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        ["PU1", "PU2"], rows, demo_attributes=_demo_attrs(),
    ))
    file_id = resp.get_json()["file_id"]

    export_resp = client.get(f"/api/ai_data_gen/export?file_id={file_id}")
    assert export_resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(export_resp.data))

    # Custom attributes are demographic metadata, not SEM indicators -- the
    # Survey Data sheet must stay exactly as before.
    data_ws = wb["Survey Data"]
    assert [c.value for c in next(data_ws.iter_rows(min_row=1, max_row=1))] == ["respondent_id", "PU1", "PU2"]

    profile_ws = wb["Respondent Profile"]
    profile_all_values = [cell.value for row in profile_ws.iter_rows() for cell in row]
    assert "Additional attribute: Income" in profile_all_values
    assert "Additional attribute: Education" in profile_all_values
    profile_header_row = next(
        r for r in profile_ws.iter_rows(min_row=1, max_row=profile_ws.max_row)
        if r[0].value == "respondent_id"
    )
    profile_headers = [c.value for c in profile_header_row]
    assert income_col in profile_headers
    assert edu_col in profile_headers

    stats_ws = wb["Descriptive Statistics"]
    stats_values = [cell.value for row in stats_ws.iter_rows() for cell in row]
    assert "Income" in stats_values
    assert "Education" in stats_values
    assert "Bachelor" in stats_values
    assert "Master" in stats_values


def test_finalize_persists_construct_theories(client):
    theories = {"Perceived Usefulness": {"citation_apa": "Davis, F. D. (1989). ...", "doi": "10.2307/249008"}}
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        ["PU1"], _rows(30, columns=("PU1",)), construct_theories=theories,
    ))
    assert resp.status_code == 200, resp.get_json()
    file_id = resp.get_json()["file_id"]
    with client.application.test_request_context():
        meta, _demo_df = ai_data_gen_api._load_ai_gen_metadata(file_id)
    assert meta["construct_theories"] == theories


def test_finalize_defaults_construct_theories_to_empty_dict(client):
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(["PU1"], _rows(30, columns=("PU1",))))
    assert resp.status_code == 200, resp.get_json()
    file_id = resp.get_json()["file_id"]
    with client.application.test_request_context():
        meta, _demo_df = ai_data_gen_api._load_ai_gen_metadata(file_id)
    assert meta["construct_theories"] == {}


def test_export_includes_references_sheet_when_construct_theories_present(client):
    theories = {"Perceived Usefulness": {"citation_apa": "Davis, F. D. (1989). ...", "doi": "10.2307/249008"}}
    resp = client.post("/api/ai_data_gen/finalize", json=_finalize_payload(
        ["PU1"], _rows(30, columns=("PU1",)), construct_theories=theories,
    ))
    file_id = resp.get_json()["file_id"]

    export_resp = client.get(f"/api/ai_data_gen/export?file_id={file_id}")
    assert export_resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(export_resp.data))
    assert wb.sheetnames == ["Survey Data", "Respondent Profile", "Descriptive Statistics", "Prompt Transparency", "References"]

    refs_ws = wb["References"]
    header = [c.value for c in next(refs_ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["Construct", "APA Citation", "DOI"]
    row = [c.value for c in next(refs_ws.iter_rows(min_row=2, max_row=2))]
    assert row[0] == "Perceived Usefulness"
    assert row[1] == "Davis, F. D. (1989). ..."
    assert row[2] == "https://doi.org/10.2307/249008"
