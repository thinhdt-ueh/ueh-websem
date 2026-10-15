"""Word export for the Power Analysis and Sample Size Sensitivity pages:
the page posts back the result it already holds (plus chart snapshots) and
gets a .docx built from it -- see pls/analysis_reports.py."""

from __future__ import annotations

import io

from docx import Document

from tests.test_api import _upload

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# 1x1 transparent PNG, standing in for a canvas.toDataURL() chart snapshot.
PNG_DATA_URL = (
    "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _power_payload(model, **extra):
    """Population parameters for every path/construct (power analysis never reads data)."""
    return {
        "model": model, "method": "cbsem", "n_mc": 20,
        "path_values": {f"{p['source']}->{p['target']}": 0.3 for p in model["paths"]},
        "loading_values": {c["id"]: 0.75 for c in model["constructs"]},
        **extra,
    }


def _docx_text(resp) -> tuple[str, Document]:
    assert resp.status_code == 200, resp.get_json(silent=True)
    assert resp.mimetype == DOCX
    doc = Document(io.BytesIO(resp.data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts), doc


def test_power_export_word(client, tam_model_json):
    run = client.post("/api/power_analysis", json=_power_payload(tam_model_json, n_from=100, n_to=200, n_step=100, lang="en"))
    assert run.status_code == 200, run.get_json()
    result = run.get_json()

    resp = client.post("/api/power_analysis/export/word", json={
        "result": result, "images": [{"label": "Power curve", "dataUrl": PNG_DATA_URL}], "lang": "en",
    })
    text, doc = _docx_text(resp)
    assert "Power Analysis Report" in text
    assert "Minimum n (power ≥ 0.80)" in text
    first = result["points"][0]
    assert f"{first['source_name']} → {first['target_name']}" in text
    assert len(doc.inline_shapes) == 1


def test_power_export_word_vietnamese_without_images(client, tam_model_json):
    run = client.post("/api/power_analysis", json=_power_payload(tam_model_json, n_from=100, n_to=150, n_step=50))
    resp = client.post("/api/power_analysis/export/word", json={"result": run.get_json(), "lang": "vi"})
    text, doc = _docx_text(resp)
    assert "Báo cáo phân tích Power" in text
    assert len(doc.inline_shapes) == 0


def test_sensitivity_export_word_shrink(client, tam_df, tam_model_json):
    file_id = _upload(client, tam_df)
    run = client.post("/api/sensitivity", json={
        "file_id": file_id, "model": tam_model_json, "lang": "en", "method": "cbsem", "step": 60,
    })
    assert run.status_code == 200, run.get_json()
    result = run.get_json()

    resp = client.post("/api/sensitivity/export/word", json={
        "result": result, "mode": "shrink", "images": [{"label": "R²", "dataUrl": PNG_DATA_URL}], "lang": "en",
    })
    text, doc = _docx_text(resp)
    assert "Sample Size Sensitivity Report" in text
    assert "R² by Sample Size" in text
    assert str(result["points"][0]["n"]) in text
    assert len(doc.inline_shapes) == 1


def test_sensitivity_export_word_resample(client, tam_df, tam_model_json):
    file_id = _upload(client, tam_df)
    run = client.post("/api/sensitivity_resample", json={
        "file_id": file_id, "model": tam_model_json, "lang": "en", "method": "cbsem",
        "new_n": 100, "n_iterations": 20,
    })
    assert run.status_code == 200, run.get_json()

    resp = client.post("/api/sensitivity/export/word", json={"result": run.get_json(), "mode": "resample", "lang": "en"})
    text, _ = _docx_text(resp)
    assert "(Resampling)" in text
    assert "Distribution of Path Coefficients" in text
    assert "Share p < 0.05" in text


def test_export_word_rejects_missing_result(client):
    for url in ("/api/power_analysis/export/word", "/api/sensitivity/export/word"):
        resp = client.post(url, json={"lang": "en"})
        assert resp.status_code == 400
        assert "error" in resp.get_json()


def test_export_word_skips_malformed_images(client, tam_model_json):
    run = client.post("/api/power_analysis", json=_power_payload(tam_model_json, n_from=100, n_to=150, n_step=50))
    resp = client.post("/api/power_analysis/export/word", json={
        "result": run.get_json(), "lang": "en",
        "images": [{"label": "bad", "dataUrl": "data:image/png;base64,!!!notpng"}, {"dataUrl": "javascript:x"}, None],
    })
    _, doc = _docx_text(resp)
    assert len(doc.inline_shapes) == 0
