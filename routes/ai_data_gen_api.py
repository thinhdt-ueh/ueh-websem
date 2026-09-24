"""AI-generated synthetic survey data for Step 1, as an alternative to
uploading a real file. Reuses the exact same multi-provider relay functions
as `routes/ai_report_api.py` (OpenAI/Gemini/Claude, caller's own API key) --
the "send one system+user message, get text back" shape is identical, only
the prompt content and the caller's handling of the response differ.

An LLM cannot reliably emit hundreds of clean numeric CSV rows in a single
completion, so generation is batched (see AI_BATCH_SIZE): the frontend calls
POST /api/ai_data_gen/batch once per chunk (each a short, independent HTTP
request -- see the module docstring in that route for why this is not one
long server-side loop), then POST /api/ai_data_gen/finalize once with all
accumulated rows. Finalize saves the assembled DataFrame into the SAME
upload directory /api/upload and /api/sample already use, under a fresh
file_id -- so the rest of the app (Step 2 model builder, /api/analyze) needs
no changes at all to treat AI-generated data exactly like an uploaded file.

Privacy: `api_key` is used only as a local variable for the one outbound
call per batch. Never written to a file, a database, or a log line.
"""

from __future__ import annotations

import io
import json
import math
import os
import re
import urllib.error
import urllib.request
import uuid
from datetime import datetime, timezone

import pandas as pd
from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook

from i18n import get_lang, t
from routes.ai_report_api import (
    DEFAULT_MODELS,
    DEFAULT_TEMPERATURE,
    MAX_TEMPERATURE,
    MIN_TEMPERATURE,
    _call_claude,
    _call_gemini,
    _call_openai,
)
from routes.api import _clean, _upload_dir
from pls.model import Model, ModelError

ai_data_gen_api = Blueprint("ai_data_gen_api", __name__, url_prefix="/api")

MAX_AI_ROWS = 500
# No minimum-row rule exists anywhere else in the app; this is just a floor
# against degenerate 1-2-row requests. The UI hint separately recommends
# 100+ observations for usable SEM estimation.
MIN_AI_ROWS = 30
# No LLM reliably emits hundreds of clean CSV rows in one completion --
# chunking keeps each response small enough to parse/validate reliably and
# gives genuine per-batch progress in the UI for free. The user can
# configure the actual rows-per-call within [MIN_BATCH_SIZE, MAX_BATCH_SIZE];
# AI_BATCH_SIZE is only the suggested default. MAX_BATCH_SIZE is enforced in
# /batch as a hard ceiling regardless of what the client requests.
# Lowered from 40: each row now also carries a `persona_description` free-
# text sentence (see PERSONA_COLUMN) on top of the indicator/demographic
# columns, and Gemini/Claude are called with a fixed output-token ceiling
# (see _call_ai_provider_mapped) shared with the AI-report feature -- a wide
# codebook (many indicator columns) at the old default of 40 rows/call could
# exceed that ceiling mid-batch, silently truncating the response short of
# the requested row count and failing the whole batch after retries.
AI_BATCH_SIZE = 25
MIN_BATCH_SIZE = 1
MAX_BATCH_SIZE = 50
# 1 initial attempt + 2 corrective retries, all within one /batch request.
# A provider HTTP error (bad key, rate limit, ...) is never retried here --
# retrying won't fix those; the frontend's own "Retry batch" button is the
# recovery path for that class of failure instead.
MAX_BATCH_ATTEMPTS = 3
LIKERT_SCALES = {5: (1, 5), 7: (1, 7)}

# Per-respondent demographic columns the AI is asked to assign alongside the
# Likert answers, kept OUT of the indicator dataset used for SEM analysis
# (see _ai_meta_dir) so age/gender never show up as candidate constructs in
# Step 2's model builder.
DEMO_COLUMNS = ["resp_age", "resp_gender"]
# The persona profile the AI must construct and role-play as before
# answering (see GEN_PERSONA_INSTRUCTION) -- forcing it to write this out
# explicitly (rather than just "silently imagine" one) measurably improves
# answer consistency. Placed FIRST in every row so the model commits to a
# persona before generating the Likert answers that follow it, matching how
# text is actually generated left-to-right.
PERSONA_COLUMN = "persona_description"
# A codebook item's "type": "likert" (default, numeric 1..N scale, becomes a
# candidate SEM indicator) or "qualitative" (open-ended free text -- kept OUT
# of the indicator dataset for the exact same reason DEMO_COLUMNS is, and
# merged into the respondent-profile file instead; see finalize()).
CODEBOOK_TYPES = {"likert", "qualitative"}
GENDER_VALUES = {"male", "female"}
DEFAULT_AGE_MIN = 18
DEFAULT_AGE_MAX = 65

# User-definable extra demographic attributes (beyond the fixed age/gender
# above), e.g. "monthly income" (numeric) or "education level" (categorical).
# Kept as a parallel, generic mechanism rather than folding age/gender into
# it, since age/gender already have dedicated, tested UI/prompt/stats code.
MAX_CUSTOM_DEMO_ATTRS = 6
MAX_CATEGORICAL_OPTIONS = 8
CUSTOM_ATTR_TYPES = {"numeric", "categorical"}


def _slugify_attr_name(name: str, existing: set[str]) -> str:
    """Column name for a custom demographic attribute. Always `demo_`-
    prefixed so it can never collide with `resp_age`/`resp_gender` or with
    an indicator column -- the frontend never needs to know this value, it
    only ever sends/receives the attribute's display `name`."""
    base = re.sub(r"[^a-z0-9]+", "_", name.strip().lower()).strip("_") or "attr"
    base = f"demo_{base}"
    candidate = base
    i = 2
    while candidate in existing:
        candidate = f"{base}_{i}"
        i += 1
    return candidate


def _validate_demo_attributes(raw, lang: str):
    """Returns (cleaned_list, None) or (None, error_message). Each cleaned
    item is {name, column, type, min, max, options} -- min/max set only for
    type == "numeric", options only for type == "categorical"."""
    if raw is None:
        return [], None
    if not isinstance(raw, list):
        return None, t("err_ai_gen_invalid_demo_attr", lang, name="?", detail="not a list")
    if len(raw) > MAX_CUSTOM_DEMO_ATTRS:
        return None, t("err_ai_gen_too_many_demo_attrs", lang, max=MAX_CUSTOM_DEMO_ATTRS)

    existing_columns: set[str] = set()
    cleaned = []
    for item in raw:
        item = item or {}
        name = str(item.get("name") or "").strip()
        attr_type = str(item.get("type") or "").strip().lower()
        if not name:
            return None, t("err_ai_gen_invalid_demo_attr", lang, name="?", detail="missing name")
        if attr_type not in CUSTOM_ATTR_TYPES:
            return None, t("err_ai_gen_invalid_demo_attr", lang, name=name, detail="invalid type")

        column = _slugify_attr_name(name, existing_columns)
        existing_columns.add(column)

        if attr_type == "numeric":
            try:
                amin = int(item.get("min"))
                amax = int(item.get("max"))
            except (TypeError, ValueError):
                return None, t("err_ai_gen_invalid_demo_attr", lang, name=name, detail="invalid min/max")
            if amin > amax:
                amin, amax = amax, amin
            cleaned.append({"name": name, "column": column, "type": "numeric", "min": amin, "max": amax, "options": None})
        else:
            raw_options = item.get("options") or []
            if not isinstance(raw_options, list):
                return None, t("err_ai_gen_invalid_demo_attr", lang, name=name, detail="invalid options")
            seen_opt = set()
            options = []
            for opt in raw_options:
                opt = str(opt or "").strip()
                if opt and opt.lower() not in seen_opt:
                    seen_opt.add(opt.lower())
                    options.append(opt)
            options = options[:MAX_CATEGORICAL_OPTIONS]
            if len(options) < 2:
                return None, t("err_ai_gen_invalid_demo_attr", lang, name=name, detail="needs at least 2 options")
            cleaned.append({"name": name, "column": column, "type": "categorical", "min": None, "max": None, "options": options})

    return cleaned, None


def _demo_attr_columns(demo_attributes: list[dict]) -> list[str]:
    return [a["column"] for a in demo_attributes]


def _normalize_categorical_value(raw: str, options: list[str]):
    """Case/whitespace-insensitive match of `raw` against the declared
    `options`, returned in the canonical declared spelling. None if no
    option matches."""
    needle = str(raw).strip().lower()
    for opt in options:
        if opt.strip().lower() == needle:
            return opt
    return None


def _ai_meta_dir() -> str:
    """Where AI-generation metadata (which respondent profile + which exact
    prompts produced a dataset) is persisted. Deliberately a SEPARATE
    subdirectory from _upload_dir()'s own files, rather than e.g.
    `file_id.json` alongside `file_id.csv` -- /api/analyze's existing
    file lookup does a prefix match (`p.startswith(file_id)`) over
    _upload_dir()'s listing, and a same-prefix metadata file there could
    shadow or race with the real data file. Keeping metadata in its own
    directory means that lookup needs no changes at all."""
    path = os.path.join(_upload_dir(), "ai_meta")
    os.makedirs(path, exist_ok=True)
    return path


def _resolve_age_bounds(demographics: dict) -> tuple[int, int]:
    def _int_or(value, default):
        try:
            return int(value)
        except (TypeError, ValueError):
            return default

    age_min = _int_or(demographics.get("age_min"), DEFAULT_AGE_MIN)
    age_max = _int_or(demographics.get("age_max"), DEFAULT_AGE_MAX)
    if age_min > age_max:
        age_min, age_max = age_max, age_min
    return age_min, age_max


def _resolve_gender_desc(demographics: dict, lang: str) -> str:
    gender = demographics.get("gender_mix") or "any"
    labels = GENDER_LABELS.get(lang, GENDER_LABELS["en"])
    return labels.get(gender, labels["any"])

GENDER_LABELS = {
    "vi": {
        "any": "không giới hạn giới tính",
        "balanced": "cân bằng nam/nữ",
        "mostly_male": "chủ yếu nam",
        "mostly_female": "chủ yếu nữ",
    },
    "en": {
        "any": "no gender constraint",
        "balanced": "a balanced mix of male/female",
        "mostly_male": "mostly male",
        "mostly_female": "mostly female",
    },
}

GEN_ROLE_FRAMING = {
    "vi": (
        "Bạn đang mô phỏng những người trả lời khảo sát thực tế cho một nghiên cứu học thuật "
        "PLS-SEM/CB-SEM. Bạn sẽ được cung cấp một bộ câu hỏi khảo sát theo thang đo Likert và "
        "mô tả đối tượng khảo sát mục tiêu."
    ),
    "en": (
        "You are simulating realistic survey respondents for an academic PLS-SEM/CB-SEM study. "
        "You will be given a codebook of Likert-scale survey questions and a target-respondent "
        "profile."
    ),
}

GEN_PERSONA_INSTRUCTION = {
    "vi": (
        "Với MỖI người trả lời (mỗi dòng), bạn phải THỰC SỰ ĐÓNG VAI người đó khi trả lời phỏng vấn "
        "-- không chỉ ngầm đoán câu trả lời hợp lý. Trước tiên, hãy DỰNG RA một hồ sơ nhân vật cụ "
        "thể, phù hợp với đối tượng mục tiêu bên dưới: độ tuổi, giới tính, nghề nghiệp/bối cảnh, và "
        "đặc biệt là thái độ/mức độ am hiểu công nghệ (VD: hài lòng hay hoài nghi, rành công nghệ "
        "hay không...). Viết hồ sơ này thành MỘT câu THẬT NGẮN GỌN (tối đa 12 từ, KHÔNG đặt tên riêng) vào "
        "cột `persona_description` -- đây là cột ĐẦU TIÊN của mỗi dòng. Sau đó, nhập vai chính xác "
        "người trong hồ sơ đó để trả lời TẤT CẢ câu hỏi còn lại -- các câu hỏi đo cùng một khái niệm "
        "(nhận biết qua nội dung câu hỏi) phải có câu trả lời tương quan hợp lý với nhau cho cùng "
        "một người, KHÔNG ngẫu nhiên độc lập từng câu. Các người trả lời khác nhau (dòng khác nhau) "
        "phải có hồ sơ và thái độ khác nhau thực sự, không chỉ là nhiễu ngẫu nhiên. TUY NHIÊN, "
        "\"tương quan hợp lý\" KHÔNG có nghĩa là giống hệt nhau: ngay cả với cùng một người, các câu "
        "hỏi đo cùng một khái niệm vẫn nên có một chút dao động tự nhiên (thường lệch nhau khoảng 1 "
        "bậc thang đo, không phải luôn luôn cùng một giá trị) -- khảo sát thật luôn có sai số đo "
        "lường, một mẫu trả lời giống hệt nhau trên mọi câu hỏi là PHI THỰC TẾ và phải tránh. LƯU Ý "
        "BẮT BUỘC: khi thêm dao động này, giá trị vẫn PHẢI nằm trong khoảng thang đo hợp lệ {lo}-{hi} "
        "-- nếu người đó có xu hướng ở sát biên (VD luôn chọn {hi}), đừng cộng thêm vượt quá {hi} hay "
        "trừ xuống dưới {lo}, hãy giữ nguyên giá trị biên đó cho câu hỏi đó thay vì đi ra ngoài khoảng "
        "cho phép. Trên "
        "toàn bộ mẫu, phong cách trả lời cũng phải trải rộng thực tế (không phải ai cũng dồn về \"đồng "
        "ý\"/\"rất đồng ý\") -- hãy có cả người trung lập và người phản đối/tiêu cực với tỉ lệ hợp lý "
        "cho đối tượng mục tiêu, và các lô sinh sau không được lặp lại xu hướng của lô trước."
    ),
    "en": (
        "For EACH respondent (row), you must ACTUALLY ROLE-PLAY as that person when answering the "
        "interview -- not just silently guess plausible answers. First, CONSTRUCT a concrete persona "
        "profile consistent with the target profile below: age, gender, occupation/context, and "
        "especially their attitude/tech-savviness (e.g., generally satisfied vs. skeptical, "
        "tech-savvy or not). Write this profile as ONE VERY SHORT sentence (12 words max, NO invented "
        "proper names) into the `persona_description` column -- this is the FIRST column of every "
        "row. Then answer ALL remaining questions fully in character as that exact person -- answers "
        "to questions measuring the same underlying concept (inferred from the question wording) "
        "must correlate realistically for that person, NOT be independently randomized. Different "
        "respondents (rows) must have genuinely different profiles and attitudes, not just "
        "item-by-item noise. HOWEVER, \"correlate realistically\" does NOT mean identical: even for "
        "the same person, items measuring the same concept should still show a little natural "
        "variation (usually about a 1-point spread on the scale, not always the exact same value) -- "
        "real surveys always have measurement error, and a response pattern that is identical across "
        "every single question is UNREALISTIC and must be avoided. MANDATORY CONSTRAINT: whenever you "
        "add this variation, the value MUST still stay within the valid scale range {lo}-{hi} -- if a "
        "person tends to sit at the edge (e.g. always picking {hi}), do NOT push a variation above "
        "{hi} or below {lo}; keep that edge value for that item instead of ever going outside the "
        "allowed range. Across the whole sample, response "
        "styles must also span a realistic range (not everyone clustering near \"agree\"/\"strongly "
        "agree\") -- include some neutral and some critical/negative respondents in a plausible "
        "proportion for the target population, and later batches must not repeat the same tendency as "
        "earlier ones."
    ),
}

# Only included (see _build_generation_messages) when the codebook has at
# least one column tagged type="qualitative" -- the codebook listing itself
# (_format_codebook) flags exactly which columns those are.
GEN_QUALITATIVE_INSTRUCTION = {
    "vi": (
        "Bộ câu hỏi bên dưới cũng có một số câu hỏi MỞ (đã đánh dấu [câu hỏi mở]), khác với các câu "
        "Likert -- với những câu này, hãy trả lời bằng 1-3 câu văn tự do, ĐÚNG với nhân vật đã dựng "
        "trong `persona_description` và nhất quán với các câu Likert cùng người đó đã trả lời (VD một "
        "người trả lời Likert thấp cho một khái niệm thì câu trả lời mở liên quan cũng nên thể hiện "
        "sự không hài lòng/hoài nghi, không mâu thuẫn). Mỗi người trả lời cần có câu trả lời mở riêng "
        "biệt, không lặp lại nguyên văn giữa các dòng."
    ),
    "en": (
        "The codebook below also has some OPEN-ENDED questions (flagged [open-ended]), unlike the "
        "Likert ones -- for those, answer in 1-3 free-form sentences, TRUE to the persona built in "
        "`persona_description` and consistent with that same person's Likert answers (e.g. someone "
        "who rated a concept low should give an open answer that also reads as dissatisfied/"
        "skeptical, not contradicting it). Each respondent needs their own distinct open answer, not "
        "reused verbatim across rows."
    ),
}

GEN_DEMO_COLUMNS_INSTRUCTION = {
    "vi": (
        "Ngoài các câu hỏi trên, mỗi người trả lời còn cần được gán CHÍNH XÁC hai thuộc tính cá "
        "nhân, khớp với hồ sơ nhân vật đã viết trong `persona_description`: tuổi (`resp_age`, số "
        "nguyên trong khoảng {age_min}-{age_max}) và giới tính (`resp_gender`, chính xác là `male` "
        "hoặc `female`). Trên toàn bộ mẫu, phân bố giới tính nên theo xu hướng: {gender_desc}."
    ),
    "en": (
        "Besides the questions above, each respondent must also be assigned exactly two personal "
        "attributes, matching the persona profile written in `persona_description`: age "
        "(`resp_age`, an integer between {age_min} and {age_max}) and gender (`resp_gender`, "
        "exactly `male` or `female`). Across the whole sample, the gender distribution should follow "
        "this tendency: {gender_desc}."
    ),
}

GEN_CUSTOM_ATTR_INSTRUCTION = {
    "vi": (
        "Ngoài ra, mỗi người trả lời còn cần được gán thêm các thuộc tính cá nhân bổ sung sau, "
        "cũng nhất quán với cá tính đã hình dung cho người đó:\n{attr_list}"
    ),
    "en": (
        "Additionally, each respondent must also be assigned the following extra personal "
        "attributes, also consistent with the persona imagined for them:\n{attr_list}"
    ),
}

GEN_OUTPUT_FORMAT = {
    "vi": (
        "Chỉ xuất ra một bảng CSV -- dòng đầu tiên là chính xác header sau: {header}. Sau đó mỗi "
        "dòng là một người trả lời: cột đầu tiên `persona_description` là hồ sơ nhân vật (một câu "
        "ngắn, LUÔN đặt trong dấu ngoặc kép \"...\" vì có thể chứa dấu phẩy), tiếp theo các câu hỏi "
        "Likert là số nguyên từ {lo} đến {hi}{qual_note}, rồi đến `resp_age` (số nguyên trong khoảng "
        "đã nêu), `resp_gender` (`male` hoặc `female`){extra_note}. KHÔNG markdown code fence, KHÔNG "
        "giải thích, KHÔNG có văn bản nào khác ngoài bảng CSV."
    ),
    "en": (
        "Output ONLY a CSV table -- the first line must be exactly this header: {header}. Each "
        "following line is one respondent: the first column `persona_description` is the persona "
        "profile (one short sentence, ALWAYS wrapped in double quotes \"...\" since it may contain "
        "commas), then the Likert questions as integers from {lo} to {hi}{qual_note}, then "
        "`resp_age` (an integer in the stated range), `resp_gender` (`male` or `female`){extra_note}. "
        "NO markdown code fences, NO explanations, NO text other than the CSV table."
    ),
}

GEN_OUTPUT_FORMAT_EXTRA_NOTE = {
    "vi": ", rồi đến các cột thuộc tính bổ sung đã khai báo ở trên, theo đúng thứ tự",
    "en": ", then the additional declared attribute columns above, in that exact order",
}

# Qualitative (open-ended) columns are interspersed among the Likert columns
# in header order (whatever order the codebook rows are in), not a separate
# block -- this note tells the model to switch format per-column rather than
# assuming the whole row is numeric.
GEN_OUTPUT_FORMAT_QUAL_NOTE = {
    "vi": (
        " (LƯU Ý: một số cột trong số đó là câu hỏi MỞ đã đánh dấu ở trên -- với các cột đó, thay vì "
        "số nguyên, hãy viết câu trả lời bằng văn bản tự do, LUÔN đặt trong dấu ngoặc kép vì có thể "
        "chứa dấu phẩy)"
    ),
    "en": (
        " (NOTE: some of those columns are the OPEN-ENDED questions flagged above -- for those "
        "columns, instead of an integer, write a free-text answer, ALWAYS wrapped in double quotes "
        "since it may contain commas)"
    ),
}

GEN_CODEBOOK_LABEL = {"vi": "Bộ câu hỏi khảo sát", "en": "Survey question codebook"}
GEN_DEMOGRAPHICS_LABEL = {"vi": "Đối tượng khảo sát mục tiêu", "en": "Target respondent profile"}
GEN_NO_DEMOGRAPHICS = {
    "vi": "Không có mô tả cụ thể về đối tượng khảo sát -- dùng đối tượng chung, đa dạng.",
    "en": "No specific target population described -- use a general, varied population.",
}


def _format_codebook(codebook: list[dict], lang: str) -> str:
    qual_tag = {"vi": " [câu hỏi mở -- trả lời bằng văn bản tự do]", "en": " [open-ended -- answer in free text]"}
    lines = []
    for item in codebook:
        q = item.get("question_text") or "(no question text provided)"
        construct = (item.get("construct") or "").strip()
        prefix = f"[{construct}] " if construct else ""
        suffix = qual_tag.get(lang, qual_tag["en"]) if item.get("type") == "qualitative" else ""
        lines.append(f"{prefix}{item['column']}: {q}{suffix}")
    return "\n".join(lines)


def _format_demographics(demographics: dict, lang: str) -> str:
    parts = []
    age_min, age_max = demographics.get("age_min"), demographics.get("age_max")
    if age_min or age_max:
        age_txt = {
            "vi": f"Độ tuổi: {age_min or '?'}-{age_max or '?'}",
            "en": f"Age range: {age_min or '?'}-{age_max or '?'}",
        }[lang]
        parts.append(age_txt)
    gender = demographics.get("gender_mix")
    labels = GENDER_LABELS.get(lang, GENDER_LABELS["en"])
    if gender in labels and gender != "any":
        gender_txt = {"vi": f"Giới tính: {labels[gender]}", "en": f"Gender: {labels[gender]}"}[lang]
        parts.append(gender_txt)
    occupation = (demographics.get("occupation") or "").strip()
    if occupation:
        parts.append({"vi": f"Nghề nghiệp/học vấn: {occupation}", "en": f"Occupation/education: {occupation}"}[lang])
    location = (demographics.get("location") or "").strip()
    if location:
        parts.append({"vi": f"Khu vực: {location}", "en": f"Location: {location}"}[lang])
    target = (demographics.get("target_population") or "").strip()
    if target:
        parts.append(target)
    if not parts:
        return GEN_NO_DEMOGRAPHICS.get(lang, GEN_NO_DEMOGRAPHICS["en"])
    return "\n".join(parts)


def _format_custom_attrs(demo_attributes: list[dict], lang: str) -> str:
    lines = []
    for attr in demo_attributes:
        if attr["type"] == "numeric":
            desc = {
                "vi": f"`{attr['column']}`: số nguyên trong khoảng {attr['min']}-{attr['max']} ({attr['name']})",
                "en": f"`{attr['column']}`: an integer between {attr['min']} and {attr['max']} ({attr['name']})",
            }[lang]
        else:
            opts = ", ".join(f"`{o}`" for o in attr["options"])
            desc = {
                "vi": f"`{attr['column']}`: đúng một trong {opts} ({attr['name']})",
                "en": f"`{attr['column']}`: exactly one of {opts} ({attr['name']})",
            }[lang]
        lines.append(f"- {desc}")
    return "\n".join(lines)


def _build_generation_messages(codebook: list[dict], demographics: dict, demo_attributes: list[dict], n_rows: int, likert_scale: int, lang: str) -> tuple[str, str]:
    lo, hi = LIKERT_SCALES[likert_scale]
    demo_attr_columns = _demo_attr_columns(demo_attributes)
    _likert_cols, qual_cols = _split_codebook_columns(codebook)
    header = ",".join([PERSONA_COLUMN] + [item["column"] for item in codebook] + DEMO_COLUMNS + demo_attr_columns)
    age_min, age_max = _resolve_age_bounds(demographics)
    gender_desc = _resolve_gender_desc(demographics, lang)
    codebook_label = GEN_CODEBOOK_LABEL.get(lang, GEN_CODEBOOK_LABEL["en"])
    demo_label = GEN_DEMOGRAPHICS_LABEL.get(lang, GEN_DEMOGRAPHICS_LABEL["en"])
    extra_note = GEN_OUTPUT_FORMAT_EXTRA_NOTE.get(lang, GEN_OUTPUT_FORMAT_EXTRA_NOTE["en"]) if demo_attributes else ""
    qual_note = GEN_OUTPUT_FORMAT_QUAL_NOTE.get(lang, GEN_OUTPUT_FORMAT_QUAL_NOTE["en"]) if qual_cols else ""
    parts = [
        GEN_ROLE_FRAMING.get(lang, GEN_ROLE_FRAMING["en"]),
        GEN_PERSONA_INSTRUCTION.get(lang, GEN_PERSONA_INSTRUCTION["en"]).format(lo=lo, hi=hi),
        f"{codebook_label}:\n{_format_codebook(codebook, lang)}",
        f"{demo_label}:\n{_format_demographics(demographics, lang)}",
        GEN_DEMO_COLUMNS_INSTRUCTION.get(lang, GEN_DEMO_COLUMNS_INSTRUCTION["en"]).format(
            age_min=age_min, age_max=age_max, gender_desc=gender_desc,
        ),
    ]
    if qual_cols:
        parts.append(GEN_QUALITATIVE_INSTRUCTION.get(lang, GEN_QUALITATIVE_INSTRUCTION["en"]))
    if demo_attributes:
        parts.append(
            GEN_CUSTOM_ATTR_INSTRUCTION.get(lang, GEN_CUSTOM_ATTR_INSTRUCTION["en"]).format(
                attr_list=_format_custom_attrs(demo_attributes, lang),
            )
        )
    parts.append(GEN_OUTPUT_FORMAT.get(lang, GEN_OUTPUT_FORMAT["en"]).format(header=header, lo=lo, hi=hi, extra_note=extra_note, qual_note=qual_note))
    system_msg = "\n\n".join(parts)
    user_msg = {
        "vi": f"Hãy sinh dữ liệu khảo sát tổng hợp cho nghiên cứu mô tả ở trên. Tổng số người trả lời cần: {n_rows}.",
        "en": f"Generate synthetic survey responses for the study described above. Total respondents needed: {n_rows}.",
    }.get(lang, f"Generate synthetic survey responses for the study described above. Total respondents needed: {n_rows}.")
    return system_msg, user_msg


def _batch_instruction(
    columns: list[str], lo: int, hi: int, start_row: int, end_row: int, lang: str,
    demo_attr_columns: list[str] | None = None, qual_columns: list[str] | None = None,
) -> str:
    n = end_row - start_row + 1
    header = ",".join([PERSONA_COLUMN] + list(columns) + DEMO_COLUMNS + (demo_attr_columns or []))
    qual_reminder = {
        "vi": f" (trừ các cột định tính {', '.join(qual_columns)} -- viết văn bản tự do, đặt trong dấu ngoặc kép, cho các cột đó)",
        "en": f" (except the qualitative columns {', '.join(qual_columns)} -- write free text, double-quoted, for those)",
    }.get(lang, f" (except the qualitative columns {', '.join(qual_columns)} -- write free text, double-quoted, for those)") if qual_columns else ""
    return {
        "vi": (
            f"Sinh chính xác {n} người trả lời ngay bây giờ (đại diện người thứ {start_row}-{end_row} "
            f"trong tổng mẫu -- đa dạng hoá cá tính so với những người đã sinh trước đó nếu có). Chỉ "
            f"xuất CSV, dòng đầu là header: {header}, theo sau đúng {n} dòng dữ liệu -- mỗi dòng bắt "
            f"đầu bằng `{PERSONA_COLUMN}` (đặt trong dấu ngoặc kép), rồi đến các câu hỏi là số nguyên "
            f"từ {lo} đến {hi}{qual_reminder}."
        ),
        "en": (
            f"Generate exactly {n} respondents now (representing respondents #{start_row}-#{end_row} "
            f"of the full sample -- vary personas from any generated before). Output ONLY the CSV, "
            f"header row: {header}, followed by exactly {n} data rows -- each row starting with "
            f"`{PERSONA_COLUMN}` (double-quoted), then whole numbers from {lo} to {hi} for the "
            f"questions{qual_reminder}."
        ),
    }.get(lang, (
        f"Generate exactly {n} respondents now (representing respondents #{start_row}-#{end_row} "
        f"of the full sample). Output ONLY the CSV, header row: {header}, followed by exactly {n} "
        f"data rows -- each row starting with `{PERSONA_COLUMN}` (double-quoted), then whole numbers "
        f"from {lo} to {hi} for the questions{qual_reminder}."
    ))


def _corrective_note(
    reason: str, columns: list[str], expected_n: int, lang: str,
    demo_attr_columns: list[str] | None = None, likert_min: int | None = None, likert_max: int | None = None,
) -> str:
    header = ",".join([PERSONA_COLUMN] + list(columns) + DEMO_COLUMNS + (demo_attr_columns or []))
    # A "got fewer rows than expected" reason means the response was cut off
    # partway (almost always an output-length limit), not a formatting
    # mistake -- the generic "do it right" note alone tends to just repeat
    # the same shortfall. Give it a concrete, actionable way to fit within
    # whatever cut it off last time: shrink the one deliberately verbose
    # field (persona_description) and stop padding/elaborating anywhere else.
    is_shortfall = reason.startswith("expected ") and "rows, got" in reason
    shortfall_hint = {
        "vi": (
            " Phản hồi trước bị CẮT NGANG giữa chừng vì quá dài. Lần này hãy viết `persona_description` "
            "cực kỳ ngắn (tối đa 6 từ) và không thêm bất kỳ văn bản dư thừa nào khác, để chắc chắn hoàn "
            "thành ĐỦ số dòng yêu cầu."
        ),
        "en": (
            " The previous response was CUT OFF partway because it ran too long. This time, write "
            "`persona_description` extremely short (6 words max) and add no other extra text anywhere, "
            "so you can actually finish the FULL requested row count."
        ),
    }.get(lang, (
        " The previous response was CUT OFF partway because it ran too long. This time, write "
        "`persona_description` extremely short (6 words max) and add no other extra text anywhere, "
        "so you can actually finish the FULL requested row count."
    )) if is_shortfall else ""
    # A range-violation is almost always the "add ~1-point natural variation"
    # instruction overshooting the boundary (e.g. 5 -> 6) rather than a
    # genuine misunderstanding of the scale -- name that exact mechanism so
    # the retry doesn't just repeat the same overshoot.
    is_range_violation = reason.startswith("a value outside the [") and likert_min is not None
    range_hint = {
        "vi": (
            f" Phản hồi trước có giá trị NẰM NGOÀI khoảng {likert_min}-{likert_max} -- rất có thể do "
            f"cộng/trừ dao động tự nhiên (~1 bậc) làm giá trị vượt quá biên. Lần này, mọi giá trị PHẢI "
            f"nằm trong {likert_min}-{likert_max}; nếu người đó đang ở giá trị biên ({likert_min} hoặc "
            f"{likert_max}), GIỮ NGUYÊN giá trị biên đó cho câu hỏi đó thay vì cộng/trừ thêm ra ngoài."
        ),
        "en": (
            f" The previous response had a value OUTSIDE the {likert_min}-{likert_max} range -- most "
            f"likely the ~1-point natural-variation instruction pushed a value past the boundary. This "
            f"time, every value MUST stay within {likert_min}-{likert_max}; if that person is already at "
            f"an edge value ({likert_min} or {likert_max}), KEEP that edge value for that item instead of "
            f"adding/subtracting past it."
        ),
    }.get(lang, (
        f" The previous response had a value OUTSIDE the {likert_min}-{likert_max} range -- most likely "
        f"the ~1-point natural-variation instruction pushed a value past the boundary. This time, every "
        f"value MUST stay within {likert_min}-{likert_max}; if that person is already at an edge value "
        f"({likert_min} or {likert_max}), KEEP that edge value for that item instead of adding/subtracting "
        f"past it."
    )) if is_range_violation else ""
    shortfall_hint = shortfall_hint + range_hint
    return {
        "vi": (
            f"Phản hồi trước không hợp lệ ({reason}). CHỈ xuất bảng CSV với header CHÍNH XÁC: "
            f"{header}, và ĐÚNG {expected_n} dòng dữ liệu -- không markdown, không giải thích.{shortfall_hint}"
        ),
        "en": (
            f"Your previous response was invalid ({reason}). Output ONLY a CSV table with header "
            f"EXACTLY: {header}, and EXACTLY {expected_n} data rows -- no markdown, no explanation.{shortfall_hint}"
        ),
    }.get(lang, (
        f"Your previous response was invalid ({reason}). Output ONLY a CSV table with header "
        f"EXACTLY: {header}, and EXACTLY {expected_n} data rows -- no markdown, no explanation.{shortfall_hint}"
    ))


def _extract_csv_block(text: str) -> str:
    stripped = text.strip()
    m = re.match(r"^```(?:csv)?\s*\n(.*?)\n```$", stripped, re.DOTALL)
    return m.group(1).strip() if m else stripped


def _parse_batch_csv(
    text: str, columns: list[str], likert_min: int, likert_max: int, expected_n: int, age_min: int, age_max: int,
    demo_attributes: list[dict] | None = None, qual_columns: list[str] | None = None,
):
    """Returns (dataframe, None) on success or (None, reason) on failure --
    never raises, so the caller's retry loop can treat every failure mode
    uniformly. The returned dataframe carries `PERSONA_COLUMN + columns +
    DEMO_COLUMNS +` any custom demographic attribute columns. `columns` is
    every codebook column (both types); `qual_columns` (a subset of it)
    marks which ones are free text instead of a Likert integer."""
    demo_attributes = demo_attributes or []
    qual_columns = qual_columns or []
    likert_columns = [c for c in columns if c not in qual_columns]
    demo_attr_columns = _demo_attr_columns(demo_attributes)
    all_columns = [PERSONA_COLUMN] + list(columns) + DEMO_COLUMNS + demo_attr_columns
    cleaned = _extract_csv_block(text)
    try:
        df = pd.read_csv(io.StringIO(cleaned))
    except Exception as exc:  # noqa: BLE001
        return None, f"could not parse as CSV ({exc})"

    if set(df.columns) != set(all_columns):
        missing = sorted(set(all_columns) - set(df.columns))
        extra = sorted(set(df.columns) - set(all_columns))
        return None, f"column mismatch (missing={missing}, extra={extra})"
    df = df[all_columns]

    if len(df) > expected_n:
        df = df.iloc[:expected_n].reset_index(drop=True)
    elif len(df) < expected_n:
        # A small shortfall (the model slightly under-counting, not a real
        # error) is accepted as a partial batch -- the caller's adaptive loop
        # simply asks for the remaining rows in a follow-up call, which is
        # far cheaper than discarding otherwise-valid rows and retrying the
        # whole batch from scratch for an exact match that isn't guaranteed
        # from any provider. Only a LARGE shortfall (more likely a genuine
        # truncation/formatting problem, see AI_BATCH_SIZE's docstring)
        # still fails and goes through the corrective-retry path below.
        if len(df) == 0 or len(df) < math.ceil(expected_n / 2):
            return None, f"expected {expected_n} rows, got {len(df)}"

    persona_raw = df[PERSONA_COLUMN]
    persona_part = persona_raw.astype(str).str.strip()
    persona_empty = persona_raw.isna() | (persona_part == "") | (persona_part.str.lower() == "nan")
    if persona_empty.any():
        return None, f"{PERSONA_COLUMN} must not be empty -- the AI must role-play as a described persona"

    likert_part = df[likert_columns].apply(pd.to_numeric, errors="coerce") if likert_columns else pd.DataFrame(index=df.index)
    if likert_part.isna().any().any():
        return None, "a non-numeric value was found"
    if ((likert_part < likert_min) | (likert_part > likert_max)).any().any():
        return None, f"a value outside the [{likert_min}, {likert_max}] range was found"

    qual_raw = df[qual_columns] if qual_columns else pd.DataFrame(index=df.index)
    qual_part = qual_raw.astype(str).apply(lambda s: s.str.strip()) if qual_columns else qual_raw
    if qual_columns:
        # astype(str) on a missing value doesn't reliably become the literal
        # string "nan" across pandas versions/dtypes (e.g. pandas 3's string
        # dtype keeps it as an actual missing marker) -- check .isna() on the
        # raw column directly too, exactly like the persona_empty check above.
        qual_empty = qual_raw.isna() | (qual_part == "") | (qual_part.apply(lambda s: s.str.lower()) == "nan")
        if qual_empty.any().any():
            return None, "a qualitative answer was empty"

    age_part = pd.to_numeric(df["resp_age"], errors="coerce")
    if age_part.isna().any():
        return None, "resp_age has a non-numeric value"
    if ((age_part < age_min) | (age_part > age_max)).any():
        return None, f"resp_age outside the [{age_min}, {age_max}] range was found"

    gender_part = df["resp_gender"].astype(str).str.strip().str.lower()
    if not gender_part.isin(GENDER_VALUES).all():
        return None, f"resp_gender must be one of {sorted(GENDER_VALUES)}"

    out = pd.DataFrame({PERSONA_COLUMN: persona_part})
    if likert_columns:
        out[likert_columns] = likert_part.astype(int)
    if qual_columns:
        out[qual_columns] = qual_part
    # Restore the original codebook column order (the two blocks above are
    # each internally ordered, but interleaved likert/qualitative columns
    # would otherwise end up likert-first, qualitative-second).
    out = out[[PERSONA_COLUMN] + list(columns)]
    out["resp_age"] = age_part.astype(int)
    out["resp_gender"] = gender_part

    for attr in demo_attributes:
        col = attr["column"]
        if attr["type"] == "numeric":
            part = pd.to_numeric(df[col], errors="coerce")
            if part.isna().any():
                return None, f"{col} has a non-numeric value"
            if ((part < attr["min"]) | (part > attr["max"])).any():
                return None, f"{col} outside the [{attr['min']}, {attr['max']}] range was found"
            out[col] = part.astype(int)
        else:
            normalized = df[col].apply(lambda v: _normalize_categorical_value(v, attr["options"]))
            if normalized.isna().any():
                return None, f"{col} must be one of {attr['options']}"
            out[col] = normalized

    return out, None


def _validate_codebook(raw_codebook, lang: str):
    if not isinstance(raw_codebook, list) or not raw_codebook:
        return None, t("err_ai_gen_missing_codebook", lang)
    seen = set()
    cleaned = []
    for item in raw_codebook:
        col = str((item or {}).get("column") or "").strip()
        if not col:
            return None, t("err_ai_gen_missing_codebook", lang)
        if col in seen:
            return None, t("err_ai_gen_duplicate_column", lang, name=col)
        seen.add(col)
        raw_type = str((item or {}).get("type") or "").strip().lower()
        cleaned.append({
            "column": col,
            "question_text": str((item or {}).get("question_text") or "").strip(),
            "construct": str((item or {}).get("construct") or "").strip(),
            "type": raw_type if raw_type in CODEBOOK_TYPES else "likert",
        })
    return cleaned, None


def _split_codebook_columns(codebook: list[dict]) -> tuple[list[str], list[str]]:
    """Returns (likert_columns, qualitative_columns), preserving codebook
    order -- the single source of truth every route derives the split from,
    so /batch, /finalize, and descriptive-stats all agree on which columns
    are SEM-indicator candidates vs. free-text respondent context."""
    likert_columns = [c["column"] for c in codebook if c.get("type") != "qualitative"]
    qual_columns = [c["column"] for c in codebook if c.get("type") == "qualitative"]
    return likert_columns, qual_columns


# ---- AI-assisted construct/indicator search (literature review) ----
# A one-shot suggestion the user reviews and selectively adds to the
# codebook table -- unlike /batch, there is no corrective-retry loop here;
# a malformed response is simply surfaced as an error the user can re-run.
MIN_CONSTRUCTS = 1
MAX_CONSTRUCTS = 10
DEFAULT_CONSTRUCTS = 5

GEN_CONSTRUCT_SEARCH_SYSTEM = {
    "vi": (
        "Bạn là trợ lý nghiên cứu hỗ trợ thiết kế bảng câu hỏi khảo sát cho một nghiên cứu "
        "PLS-SEM/CB-SEM. Dựa trên các lý thuyết/mô hình học thuật đã được công nhận (TAM, UTAUT, "
        "TPB, TRA, D&M IS Success, v.v. -- chọn mô hình phù hợp nhất với chủ đề bên dưới), hãy xác "
        "định {n_constructs} biến tiềm ẩn (construct) liên quan, mỗi construct gồm 3-5 biến quan sát "
        "(indicator) được phỏng theo các thang đo đã được kiểm định trong các nghiên cứu trước. Với "
        "mỗi construct, hãy đặt một tiền tố mã ngắn hợp lệ (chỉ chữ cái/số, VD PU, PEOU, ATT, INT) và "
        "đánh số các item lần lượt (PU1, PU2, PU3...). Với mỗi construct, hãy cung cấp thêm trích dẫn "
        "APA (phong cách APA 7) cho nguồn gốc học thuật của construct đó (tác giả, năm, tên bài "
        "báo/mô hình gốc, tạp chí) -- nếu biết DOI thì cung cấp, còn nếu KHÔNG CHẮC CHẮN thì để "
        "chuỗi rỗng thay vì bịa ra DOI giả. CHỈ xuất ra JSON hợp lệ, không dùng markdown code fence, "
        "không giải thích thêm, đúng theo cấu trúc sau: "
        '{{"constructs": [{{"name": "<tên construct>", "theory": {{"citation_apa": '
        '"<trích dẫn APA 7 đầy đủ>", "doi": "<DOI hoặc chuỗi rỗng>"}}, "items": [{{"column": '
        '"<mã ngắn>", "question_text": "<nội dung câu hỏi khảo sát đầy đủ, bằng tiếng Việt>"}}]}}]}}'
    ),
    "en": (
        "You are a research assistant helping design a survey questionnaire for a PLS-SEM/CB-SEM "
        "study. Based on established academic theories/models (TAM, UTAUT, TPB, TRA, D&M IS "
        "Success, etc. -- pick whichever best fits the topic below), identify {n_constructs} "
        "relevant latent constructs, each with 3-5 measurement items (indicators) adapted from "
        "validated scales used in prior research. For each construct, invent a short valid code "
        "prefix (letters/digits only, e.g. PU, PEOU, ATT, INT) and number its items sequentially "
        "(PU1, PU2, PU3...). For each construct, also provide an APA (7th edition) citation for its "
        "academic origin (author(s), year, original paper/model title, journal) -- provide the DOI "
        "if you know it, but if you are NOT CERTAIN, leave it as an empty string rather than "
        "inventing a fake DOI. Output ONLY valid JSON, no markdown code fence, no explanation, in "
        "exactly this shape: "
        '{{"constructs": [{{"name": "<construct name>", "theory": {{"citation_apa": '
        '"<full APA 7 citation>", "doi": "<DOI or empty string>"}}, "items": [{{"column": '
        '"<short code>", "question_text": "<full survey item wording, in English>"}}]}}]}}'
    ),
}


def _build_construct_search_messages(topic: str, n_constructs: int, lang: str) -> tuple[str, str]:
    system_msg = GEN_CONSTRUCT_SEARCH_SYSTEM.get(lang, GEN_CONSTRUCT_SEARCH_SYSTEM["en"]).format(n_constructs=n_constructs)
    user_msg = {
        "vi": f"Chủ đề / bối cảnh nghiên cứu: {topic}",
        "en": f"Research topic / context: {topic}",
    }.get(lang, f"Research topic / context: {topic}")
    return system_msg, user_msg


def _extract_json_block(text: str) -> str:
    stripped = text.strip()
    m = re.match(r"^```(?:json)?\s*\n(.*?)\n```$", stripped, re.DOTALL)
    return m.group(1).strip() if m else stripped


def _parse_construct_search_json(text: str):
    """Returns (constructs, None) on success or (None, reason) on failure --
    never raises. constructs: [{"name", "theory": {"citation_apa", "doi"},
    "items": [{"column", "question_text"}]}] with globally-unique,
    auto-suffixed column codes."""
    cleaned = _extract_json_block(text)
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        return None, f"could not parse as JSON ({exc})"

    if not isinstance(data, dict) or not isinstance(data.get("constructs"), list) or not data["constructs"]:
        return None, "missing or empty 'constructs' list"

    seen_columns: set[str] = set()
    constructs = []
    for group in data["constructs"]:
        name = str((group or {}).get("name") or "").strip()
        raw_items = (group or {}).get("items")
        raw_theory = (group or {}).get("theory") or {}
        citation_apa = str(raw_theory.get("citation_apa") or "").strip()
        doi = str(raw_theory.get("doi") or "").strip()
        if not name or not isinstance(raw_items, list) or not raw_items:
            return None, "each construct needs a non-empty 'name' and non-empty 'items'"
        if not citation_apa:
            return None, "each construct needs a non-empty 'theory.citation_apa'"
        items = []
        for raw_item in raw_items:
            column = str((raw_item or {}).get("column") or "").strip()
            question_text = str((raw_item or {}).get("question_text") or "").strip()
            if not column or not question_text:
                return None, "each item needs a non-empty 'column' and 'question_text'"
            # Auto-suffix rather than fail -- this is a one-shot suggestion
            # the user reviews before anything is committed to the codebook.
            candidate = column
            i = 2
            while candidate in seen_columns:
                candidate = f"{column}_{i}"
                i += 1
            seen_columns.add(candidate)
            items.append({"column": candidate, "question_text": question_text})
        constructs.append({"name": name, "theory": {"citation_apa": citation_apa, "doi": doi}, "items": items})
    return constructs, None


@ai_data_gen_api.post("/ai_data_gen/suggest_constructs")
def suggest_constructs():
    payload = request.get_json(force=True, silent=True) or {}
    lang = get_lang(payload)
    provider = (payload.get("provider") or "openai").strip().lower()
    api_key = (payload.get("api_key") or "").strip()
    model = (payload.get("model") or DEFAULT_MODELS.get(provider, DEFAULT_MODELS["openai"])).strip()
    topic = (payload.get("topic") or "").strip()
    try:
        temperature = float(payload.get("temperature", DEFAULT_TEMPERATURE))
    except (TypeError, ValueError):
        temperature = DEFAULT_TEMPERATURE
    temperature = max(MIN_TEMPERATURE, min(MAX_TEMPERATURE, temperature))
    try:
        n_constructs = int(payload.get("n_constructs", DEFAULT_CONSTRUCTS))
    except (TypeError, ValueError):
        n_constructs = DEFAULT_CONSTRUCTS
    n_constructs = max(MIN_CONSTRUCTS, min(MAX_CONSTRUCTS, n_constructs))

    if provider not in DEFAULT_MODELS:
        return jsonify(error=t("err_ai_bad_provider", lang)), 400
    if not api_key:
        return jsonify(error=t("err_ai_missing_key", lang)), 400
    if not topic:
        return jsonify(error=t("err_ai_gen_missing_topic", lang)), 400

    system_msg, user_msg = _build_construct_search_messages(topic, n_constructs, lang)
    text, err_response = _call_ai_provider_mapped(provider, api_key, model, system_msg, user_msg, temperature, lang)
    if err_response is not None:
        return err_response

    constructs, reason = _parse_construct_search_json(text)
    if constructs is None:
        return jsonify(error=t("err_ai_gen_bad_construct_search", lang, detail=reason)), 422
    return jsonify(constructs=constructs)


# ---- AI-assisted structural model drawing (Step 2) ----
# Same one-shot contract as suggest_constructs above -- no corrective-retry
# loop; a malformed response just 422s and the user re-runs from the modal.
MIN_CONSTRUCTS_FOR_PATHS = 2  # matches Model.from_json's own err_model_min_constructs

GEN_PATHS_SEARCH_SYSTEM = {
    "vi": (
        "Bạn là trợ lý nghiên cứu SEM (PLS-SEM/CB-SEM) am hiểu lý thuyết. Dưới đây là "
        "danh sách các biến tiềm ẩn (construct) đã có, mỗi construct kèm tên và các biến "
        "quan sát (indicator) của nó -- dựa vào tên construct và nội dung indicator, hãy "
        "suy luận vai trò lý thuyết hợp lý nhất của từng construct (tiền đề/trung gian/"
        "kết quả) và đề xuất một mô hình cấu trúc (structural model) hợp lý, PHI CHU TRÌNH "
        "(không có vòng lặp nhân quả), gồm các đường dẫn (path) một chiều giữa các "
        "construct, dựa trên logic lý thuyết đã được công nhận (VD chuỗi TAM/UTAUT-kiểu "
        "PEOU->PU->Attitude->Intention, quan hệ trung gian, v.v.) phù hợp nhất với các "
        "construct đã cho. Nếu dựa trên tên và nội dung indicator, bạn thấy một construct "
        "nào đó có khả năng đóng vai trò BIẾN ĐIỀU TIẾT (moderator) cho một mối quan hệ nào "
        "đó thay vì chỉ là một mắt xích trung gian trong chuỗi chính, hãy liệt kê construct "
        "đó vào mảng \"moderator_suggestions\" (kèm lý do ngắn gọn) thay vì cố tạo đường dẫn "
        "cho vai trò đó -- KHÔNG thể hiện vai trò điều tiết trong danh sách \"paths\" (mỗi "
        "path vẫn chỉ là một đường dẫn trực tiếp một chiều giữa hai construct); nếu không có "
        "construct nào phù hợp, trả về mảng rỗng.{interaction_note} Mỗi construct nên có ít "
        "nhất một liên kết vào mạng lưới nếu hợp lý. TUYỆT ĐỐI không tạo chu trình (A->B->A "
        "hoặc dài hơn). QUAN TRỌNG: trong \"source\"/\"target\"/\"construct_id\", PHẢI dùng "
        "CHÍNH XÁC chuỗi id xuất hiện sau \"id=\" của từng construct bên dưới (copy y nguyên "
        "từng ký tự) -- TUYỆT ĐỐI KHÔNG dùng tên construct ở đó. CHỈ xuất ra JSON hợp lệ, "
        "không dùng markdown code fence, không giải thích thêm, đúng cấu trúc sau: "
        '{{"paths": [{{"source": "<id construct>", "target": "<id construct>"}}], '
        '"moderator_suggestions": [{{"construct_id": "<id construct>", "reason": "<lý do '
        'ngắn gọn bằng tiếng Việt>"}}], '
        '"rationale": "<đoạn văn xuôi giải thích logic lý thuyết vì sao chọn mô hình này, '
        'bằng tiếng Việt, KHÔNG dùng markdown>"}}'
    ),
    "en": (
        "You are a theory-grounded SEM (PLS-SEM/CB-SEM) research assistant. Below is the "
        "list of existing latent constructs, each with its name and measurement indicators "
        "-- based on each construct's name and indicator wording, infer its most plausible "
        "theoretical role (antecedent/mediator/outcome) and propose a plausible, ACYCLIC "
        "structural model: one-directional paths between constructs, grounded in "
        "established theoretical logic (e.g. a TAM/UTAUT-style PEOU->PU->Attitude->"
        "Intention chain, mediation, etc.) that best fits the given constructs. If, based on "
        "a construct's name and indicator wording, you judge it more likely to play a "
        "MODERATING role for some relationship rather than just being another link in the "
        "main chain, list that construct in the \"moderator_suggestions\" array (with a "
        "brief reason) instead of forcing it into a path for that role -- do NOT try to "
        "represent a moderating role in the \"paths\" list itself (each path is still just a "
        "direct one-directional link between two constructs); return an empty array if none "
        "apply.{interaction_note} Every construct should have at least one connection into "
        "the network where sensible. NEVER create a cycle (A->B->A or longer). IMPORTANT: "
        "in \"source\"/\"target\"/\"construct_id\", you MUST use EXACTLY the id string shown "
        "after \"id=\" for each construct below (copy it character-for-character) -- NEVER "
        "put the construct's name there. Output ONLY valid JSON, no markdown code fence, no "
        "explanation, in exactly this shape: "
        '{{"paths": [{{"source": "<construct id>", "target": "<construct id>"}}], '
        '"moderator_suggestions": [{{"construct_id": "<construct id>", "reason": "<brief '
        'reason in English>"}}], '
        '"rationale": "<a prose paragraph explaining the theoretical logic behind this '
        'model, in English, NO markdown>"}}'
    ),
}

GEN_PATHS_INTERACTION_NOTE = {
    "vi": (
        " LƯU Ý về construct tương tác/điều tiết (đã liệt kê nguồn của nó bên dưới): nó CHỈ "
        "được là NGUỒN của một đường dẫn, KHÔNG BAO GIỜ là ĐÍCH; và bất kỳ construct nào nó "
        "trỏ tới cũng PHẢI nhận thêm đường dẫn trực tiếp từ CẢ HAI construct nguồn của nó "
        "(hiệu ứng chính/main effect), nếu không mô hình sẽ không hợp lệ."
    ),
    "en": (
        " NOTE on any interaction/moderation construct (its two source constructs are "
        "listed below it): it can ONLY be a path SOURCE, NEVER a target; and whatever "
        "construct it points to MUST also receive a direct path from BOTH of its own "
        "source constructs (the main effects), or the model will be rejected as invalid."
    ),
}


def _format_constructs_for_paths(constructs: list[dict], indicator_descriptions: dict, lang: str) -> tuple[str, bool]:
    """Returns (formatted description, has_interaction) for the given raw
    construct dicts (same shape /api/analyze accepts). indicator_descriptions
    maps indicator column -> question wording (from the AI Lab codebook, if
    any) so the AI can reason about each indicator's actual scale content,
    not just its column name."""
    by_id = {str(c.get("id", "")).strip(): c for c in constructs}
    indicator_descriptions = indicator_descriptions or {}
    lines = []
    has_interaction = False
    for c in constructs:
        cid = str(c.get("id", "")).strip()
        name = str(c.get("name", "")).strip()
        mode = str(c.get("mode", "A")).strip().upper()
        indicators = [str(i).strip() for i in (c.get("indicators") or []) if str(i).strip()]
        if mode == "I":
            has_interaction = True
            raw_pair = [str(x).strip() for x in (c.get("interaction_of") or [])]
            pair_names = [by_id[p]["name"] for p in raw_pair if p in by_id]
            src_desc = " x ".join(pair_names) if pair_names else "?"
            lines.append(f"- id={cid}, name=\"{name}\" (interaction/moderation of: {src_desc})")
        else:
            if indicators:
                ind_parts = []
                for col in indicators:
                    desc = str(indicator_descriptions.get(col, "")).strip()
                    ind_parts.append(f'{col} ("{desc}")' if desc else col)
                ind_desc = ", ".join(ind_parts)
            else:
                ind_desc = "(no indicators)"
            lines.append(f"- id={cid}, name=\"{name}\", indicators: {ind_desc}")
    return "\n".join(lines), has_interaction


def _build_paths_search_messages(
    constructs: list[dict], extra_context: str, indicator_descriptions: dict, lang: str
) -> tuple[str, str]:
    constructs_desc, has_interaction = _format_constructs_for_paths(constructs, indicator_descriptions, lang)
    interaction_note = GEN_PATHS_INTERACTION_NOTE.get(lang, GEN_PATHS_INTERACTION_NOTE["en"]) if has_interaction else ""
    system_msg = GEN_PATHS_SEARCH_SYSTEM.get(lang, GEN_PATHS_SEARCH_SYSTEM["en"]).format(interaction_note=interaction_note)
    system_msg += "\n\n" + {
        "vi": "Danh sách construct:",
        "en": "Constructs:",
    }.get(lang, "Constructs:") + "\n" + constructs_desc
    extra_context = (extra_context or "").strip()
    user_msg = {
        "vi": "Hãy đề xuất mô hình cấu trúc cho các construct trên.",
        "en": "Please propose a structural model for the constructs above.",
    }.get(lang, "Please propose a structural model for the constructs above.")
    if extra_context:
        user_msg += "\n\n" + {
            "vi": f"Bối cảnh nghiên cứu bổ sung: {extra_context}",
            "en": f"Additional research context: {extra_context}",
        }.get(lang, f"Additional research context: {extra_context}")
    return system_msg, user_msg


def _resolve_construct_refs(paths: list[dict], moderator_suggestions: list[dict], constructs: list[dict]):
    """Despite the prompt insisting on the exact id string, a model sometimes
    echoes a construct's NAME (or a case/whitespace-mangled id) in
    source/target/construct_id instead -- resolve those against the known
    constructs (exact id match first, then case-insensitive name match)
    before validating, rather than hard-failing an otherwise-correct
    proposal. Anything that still doesn't resolve is left as-is, so
    Model.from_json's "unknown construct" error still fires for a genuinely
    invented id."""
    by_id = set()
    by_name = {}
    for c in constructs:
        cid = str(c.get("id", "")).strip()
        name = str(c.get("name", "")).strip()
        if cid:
            by_id.add(cid)
        if name:
            by_name.setdefault(name.lower(), cid)

    def resolve(raw: str) -> str:
        raw = str(raw or "").strip()
        if raw in by_id:
            return raw
        return by_name.get(raw.lower(), raw)

    resolved_paths = [{"source": resolve(p["source"]), "target": resolve(p["target"])} for p in paths]
    resolved_suggestions = [
        {"construct_id": resolve(m["construct_id"]), "reason": m.get("reason", "")} for m in moderator_suggestions
    ]
    return resolved_paths, resolved_suggestions


def _parse_paths_search_json(text: str):
    """Returns (paths, rationale, moderator_suggestions, None) on success or
    (None, None, None, reason) on structural failure -- never raises.
    Cycle/interaction/self-loop/unknown-construct rules for `paths` are NOT
    re-checked here; the caller validates them via Model.from_json, reusing
    that real logic. moderator_suggestions entries are only lightly shaped
    here -- the caller drops any referring to an unknown or already-
    interaction construct id, since that needs the construct list."""
    cleaned = _extract_json_block(text)
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        return None, None, None, f"could not parse as JSON ({exc})"

    if not isinstance(data, dict):
        return None, None, None, "response was not a JSON object"
    raw_paths = data.get("paths")
    rationale = str(data.get("rationale") or "").strip()
    if not isinstance(raw_paths, list) or not raw_paths:
        return None, None, None, "missing or empty 'paths' list"
    if not rationale:
        return None, None, None, "missing or empty 'rationale'"

    paths = []
    for raw_path in raw_paths:
        source = str((raw_path or {}).get("source") or "").strip()
        target = str((raw_path or {}).get("target") or "").strip()
        if not source or not target:
            return None, None, None, "each path needs a non-empty 'source' and 'target'"
        paths.append({"source": source, "target": target})

    moderator_suggestions = []
    for raw in data.get("moderator_suggestions") or []:
        construct_id = str((raw or {}).get("construct_id") or "").strip()
        if not construct_id:
            continue
        reason = str((raw or {}).get("reason") or "").strip()
        moderator_suggestions.append({"construct_id": construct_id, "reason": reason})

    return paths, rationale, moderator_suggestions, None


@ai_data_gen_api.post("/ai_data_gen/suggest_paths_prompt")
def suggest_paths_prompt():
    """Builds the system/user prompt text for the AI-drawn-model feature
    without calling any AI provider, so the frontend can show it to the
    user for review/editing before /suggest_paths actually sends it --
    mirrors suggest_prompt's relationship to /batch."""
    payload = request.get_json(force=True, silent=True) or {}
    lang = get_lang(payload)
    constructs = payload.get("constructs") or []
    extra_context = payload.get("extra_context") or ""
    indicator_descriptions = payload.get("indicator_descriptions") or {}

    if not isinstance(constructs, list) or len(constructs) < MIN_CONSTRUCTS_FOR_PATHS:
        return jsonify(error=t("err_ai_gen_missing_constructs", lang, min=MIN_CONSTRUCTS_FOR_PATHS)), 400
    if not isinstance(indicator_descriptions, dict):
        indicator_descriptions = {}

    system_msg, user_msg = _build_paths_search_messages(constructs, extra_context, indicator_descriptions, lang)
    return jsonify(system_prompt=system_msg, user_prompt=user_msg)


@ai_data_gen_api.post("/ai_data_gen/suggest_paths")
def suggest_paths():
    payload = request.get_json(force=True, silent=True) or {}
    lang = get_lang(payload)
    provider = (payload.get("provider") or "openai").strip().lower()
    api_key = (payload.get("api_key") or "").strip()
    model = (payload.get("model") or DEFAULT_MODELS.get(provider, DEFAULT_MODELS["openai"])).strip()
    constructs = payload.get("constructs") or []
    system_msg = (payload.get("system_prompt") or "").strip()
    user_msg = (payload.get("user_prompt") or "").strip()
    try:
        temperature = float(payload.get("temperature", DEFAULT_TEMPERATURE))
    except (TypeError, ValueError):
        temperature = DEFAULT_TEMPERATURE
    temperature = max(MIN_TEMPERATURE, min(MAX_TEMPERATURE, temperature))

    if provider not in DEFAULT_MODELS:
        return jsonify(error=t("err_ai_bad_provider", lang)), 400
    if not api_key:
        return jsonify(error=t("err_ai_missing_key", lang)), 400
    if not isinstance(constructs, list) or len(constructs) < MIN_CONSTRUCTS_FOR_PATHS:
        return jsonify(error=t("err_ai_gen_missing_constructs", lang, min=MIN_CONSTRUCTS_FOR_PATHS)), 400
    if not system_msg or not user_msg:
        return jsonify(error=t("err_ai_missing_prompt", lang)), 400

    text, err_response = _call_ai_provider_mapped(provider, api_key, model, system_msg, user_msg, temperature, lang)
    if err_response is not None:
        return err_response

    paths, rationale, moderator_suggestions, reason = _parse_paths_search_json(text)
    if paths is None:
        return jsonify(error=t("err_ai_gen_bad_paths", lang, detail=reason)), 422
    paths, moderator_suggestions = _resolve_construct_refs(paths, moderator_suggestions, constructs)

    # Reuse the real structural-validity rules (cycles, interaction-target/
    # main-effect requirements, self-loops, unknown constructs) instead of
    # re-implementing any of them -- Model.from_json already enforces every
    # one of them for /api/analyze.
    try:
        Model.from_json({"constructs": constructs, "paths": paths}, lang=lang)
    except ModelError as exc:
        return jsonify(error=t("err_ai_gen_bad_paths", lang, detail=str(exc))), 422

    # Drop suggestions for a construct id the AI made up, or one that's
    # already an interaction/moderation construct (nothing actionable left
    # to do with those in the UI).
    known_ids = {str(c.get("id", "")).strip() for c in constructs}
    interaction_ids = {
        str(c.get("id", "")).strip() for c in constructs if str(c.get("mode", "A")).strip().upper() == "I"
    }
    moderator_suggestions = [
        m for m in moderator_suggestions if m["construct_id"] in known_ids and m["construct_id"] not in interaction_ids
    ]

    return jsonify(paths=paths, rationale=rationale, moderator_suggestions=moderator_suggestions)


@ai_data_gen_api.post("/ai_data_gen/suggest_prompt")
def suggest_prompt():
    payload = request.get_json(force=True, silent=True) or {}
    lang = get_lang(payload)
    codebook, err = _validate_codebook(payload.get("codebook"), lang)
    if err:
        return jsonify(error=err), 400

    demographics = payload.get("demographics") or {}
    demo_attributes, attr_err = _validate_demo_attributes(payload.get("demo_attributes"), lang)
    if attr_err:
        return jsonify(error=attr_err), 400
    try:
        n_rows = int(payload.get("n_rows", 100))
    except (TypeError, ValueError):
        n_rows = 100
    n_rows = max(MIN_AI_ROWS, min(MAX_AI_ROWS, n_rows))
    try:
        likert_scale = int(payload.get("likert_scale", 5))
    except (TypeError, ValueError):
        likert_scale = 5
    if likert_scale not in LIKERT_SCALES:
        return jsonify(error=t("err_ai_gen_invalid_likert", lang)), 400
    try:
        requested_batch_size = int(payload.get("batch_size", AI_BATCH_SIZE))
    except (TypeError, ValueError):
        requested_batch_size = AI_BATCH_SIZE
    requested_batch_size = max(MIN_BATCH_SIZE, min(MAX_BATCH_SIZE, requested_batch_size))

    system_msg, user_msg = _build_generation_messages(codebook, demographics, demo_attributes, n_rows, likert_scale, lang)
    batch_size = min(requested_batch_size, n_rows)
    total_batches = math.ceil(n_rows / batch_size)

    # Exact preview of what the first /batch call will actually append to
    # the user prompt above -- built via the same helper /batch itself uses,
    # so this is a genuine WYSIWYG preview, not an approximation.
    lo, hi = LIKERT_SCALES[likert_scale]
    _qual_columns = _split_codebook_columns(codebook)[1]
    first_batch_instruction = _batch_instruction(
        [item["column"] for item in codebook], lo, hi, 1, batch_size, lang, _demo_attr_columns(demo_attributes), _qual_columns,
    )
    return jsonify(
        system_prompt=system_msg,
        user_prompt=user_msg,
        batch_size=batch_size,
        total_batches=total_batches,
        first_batch_instruction=first_batch_instruction,
    )


def _call_ai_provider_mapped(provider: str, api_key: str, model: str, system_msg: str, user_msg: str, temperature: float, lang: str):
    """Calls the given provider and maps any failure to a Flask error
    response tuple, shared by every route in this file that makes a live
    AI call. Returns (text, None) on success or (None, (response, status))
    on failure -- never raises."""
    try:
        if provider == "openai":
            text = _call_openai(api_key, model, system_msg, user_msg, temperature)
        elif provider == "gemini":
            text = _call_gemini(api_key, model, system_msg, user_msg, temperature)
        else:
            text = _call_claude(api_key, model, system_msg, user_msg, temperature)
    except urllib.error.HTTPError as exc:
        try:
            detail = json.loads(exc.read().decode("utf-8")).get("error", {}).get("message")
        except Exception:  # noqa: BLE001
            detail = None
        if exc.code in (401, 403):
            return None, (jsonify(error=t("err_ai_invalid_key", lang)), 400)
        if exc.code in (429, 529):
            return None, (jsonify(error=t("err_ai_rate_limited", lang)), 429)
        if exc.code == 400:
            return None, (jsonify(error=t("err_ai_bad_model", lang, detail=detail or exc.reason)), 400)
        return None, (jsonify(error=t("err_ai_request_failed", lang, detail=detail or exc.reason)), 502)
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        return None, (jsonify(error=t("err_ai_request_failed", lang, detail=str(exc))), 502)
    except (KeyError, IndexError, json.JSONDecodeError) as exc:
        return None, (jsonify(error=t("err_ai_request_failed", lang, detail=str(exc))), 502)
    return text, None


@ai_data_gen_api.post("/ai_data_gen/batch")
def generate_batch():
    payload = request.get_json(force=True, silent=True) or {}
    lang = get_lang(payload)
    provider = (payload.get("provider") or "openai").strip().lower()
    api_key = (payload.get("api_key") or "").strip()
    model = (payload.get("model") or DEFAULT_MODELS.get(provider, DEFAULT_MODELS["openai"])).strip()
    system_msg = payload.get("system_prompt") or ""
    user_msg = payload.get("user_prompt") or ""
    columns = payload.get("columns") or []
    try:
        temperature = float(payload.get("temperature", DEFAULT_TEMPERATURE))
    except (TypeError, ValueError):
        temperature = DEFAULT_TEMPERATURE
    temperature = max(MIN_TEMPERATURE, min(MAX_TEMPERATURE, temperature))

    if provider not in DEFAULT_MODELS:
        return jsonify(error=t("err_ai_bad_provider", lang)), 400
    if not api_key:
        return jsonify(error=t("err_ai_missing_key", lang)), 400
    if not columns or not isinstance(columns, list):
        return jsonify(error=t("err_ai_gen_missing_codebook", lang)), 400
    demo_attributes, attr_err = _validate_demo_attributes(payload.get("demo_attributes"), lang)
    if attr_err:
        return jsonify(error=attr_err), 400
    try:
        likert_min = int(payload.get("likert_min"))
        likert_max = int(payload.get("likert_max"))
        start_row = int(payload.get("start_row"))
        end_row = int(payload.get("end_row"))
    except (TypeError, ValueError):
        return jsonify(error=t("err_ai_gen_invalid_n_rows", lang, min=MIN_AI_ROWS, max=MAX_AI_ROWS)), 400

    expected_n = end_row - start_row + 1
    if expected_n <= 0 or expected_n > MAX_BATCH_SIZE:
        return jsonify(error=t("err_ai_gen_invalid_n_rows", lang, min=MIN_AI_ROWS, max=MAX_AI_ROWS)), 400

    age_min, age_max = _resolve_age_bounds({"age_min": payload.get("demo_age_min"), "age_max": payload.get("demo_age_max")})
    demo_attr_columns = _demo_attr_columns(demo_attributes)
    categorical_attr_columns = {a["column"] for a in demo_attributes if a["type"] == "categorical"}
    raw_qual_columns = payload.get("qualitative_columns") or []
    qual_columns = [c for c in raw_qual_columns if c in columns] if isinstance(raw_qual_columns, list) else []

    batch_user_msg = user_msg + "\n\n" + _batch_instruction(columns, likert_min, likert_max, start_row, end_row, lang, demo_attr_columns, qual_columns)
    current_system_msg = system_msg
    last_reason = None

    for _attempt in range(MAX_BATCH_ATTEMPTS):
        text, err_response = _call_ai_provider_mapped(provider, api_key, model, current_system_msg, batch_user_msg, temperature, lang)
        if err_response is not None:
            return err_response

        df, reason = _parse_batch_csv(text, columns, likert_min, likert_max, expected_n, age_min, age_max, demo_attributes, qual_columns)
        if df is not None:
            # pandas' int64/object dtypes aren't natively JSON-serializable
            # (numpy int64 in particular) -- cast every cell to a plain
            # Python int/str explicitly rather than relying on jsonify.
            string_cols = {"resp_gender", PERSONA_COLUMN} | categorical_attr_columns | set(qual_columns)
            safe_rows = [
                {k: (str(v) if k in string_cols else int(v)) for k, v in row.items()}
                for row in df.to_dict(orient="records")
            ]
            return jsonify(rows=safe_rows, used_system_prompt=current_system_msg, used_user_prompt=batch_user_msg)
        last_reason = reason
        current_system_msg = system_msg + "\n\n" + _corrective_note(reason, columns, expected_n, lang, demo_attr_columns, likert_min, likert_max)

    return jsonify(error=t("err_ai_gen_bad_batch", lang, detail=last_reason)), 422


def _numeric_stats(series: pd.Series) -> dict:
    return {
        "mean": _clean(round(float(series.mean()), 3)),
        "std": _clean(round(float(series.std()), 3)) if len(series) > 1 else 0.0,
        "min": _clean(float(series.min())),
        "max": _clean(float(series.max())),
    }


def _category_stats(series: pd.Series, options: list[str], n: int) -> dict:
    counts = series.value_counts()
    return {
        opt: {"count": int(counts.get(opt, 0)), "pct": _clean(round(100 * counts.get(opt, 0) / n, 1)) if n else 0.0}
        for opt in options
    }


def _compute_descriptive_stats(indicator_df: pd.DataFrame, demo_df: pd.DataFrame, demo_attributes: list[dict] | None = None) -> dict:
    demo_attributes = demo_attributes or []
    indicators = {col: _numeric_stats(indicator_df[col]) for col in indicator_df.columns}
    n = len(demo_df)
    gender_stats = _category_stats(demo_df["resp_gender"], sorted(GENDER_VALUES), n)
    custom = []
    for attr in demo_attributes:
        if attr["column"] not in demo_df.columns:
            continue
        if attr["type"] == "numeric":
            custom.append({"name": attr["name"], "column": attr["column"], "type": "numeric", "stats": _numeric_stats(demo_df[attr["column"]])})
        else:
            custom.append({
                "name": attr["name"],
                "column": attr["column"],
                "type": "categorical",
                "counts": _category_stats(demo_df[attr["column"]], attr["options"], n),
            })
    demographics = {"age": _numeric_stats(demo_df["resp_age"]), "gender": gender_stats, "custom": custom}
    return {"indicators": indicators, "demographics": demographics}


def _save_ai_gen_metadata(file_id: str, demo_df: pd.DataFrame, meta: dict) -> None:
    meta_dir = _ai_meta_dir()
    demo_df.to_csv(os.path.join(meta_dir, f"{file_id}_respondents.csv"), index=False)
    with open(os.path.join(meta_dir, f"{file_id}.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)


def _load_ai_gen_metadata(file_id: str):
    """Returns (meta_dict, respondents_df) or (None, None) if this file_id
    was never AI-generated (e.g. a plain upload) or its metadata is gone."""
    meta_dir = _ai_meta_dir()
    meta_path = os.path.join(meta_dir, f"{file_id}.json")
    respondents_path = os.path.join(meta_dir, f"{file_id}_respondents.csv")
    if not os.path.exists(meta_path) or not os.path.exists(respondents_path):
        return None, None
    with open(meta_path, "r", encoding="utf-8") as f:
        meta = json.load(f)
    respondents_df = pd.read_csv(respondents_path)
    return meta, respondents_df


@ai_data_gen_api.post("/ai_data_gen/finalize")
def finalize():
    payload = request.get_json(force=True, silent=True) or {}
    lang = get_lang(payload)
    columns = payload.get("columns") or []
    rows = payload.get("rows") or []
    filename = (payload.get("filename") or "ai_generated_survey.csv").strip()

    if not columns or not isinstance(columns, list):
        return jsonify(error=t("err_ai_gen_missing_codebook", lang)), 400
    if not isinstance(rows, list) or not (MIN_AI_ROWS <= len(rows) <= MAX_AI_ROWS):
        return jsonify(error=t("err_ai_gen_invalid_n_rows", lang, min=MIN_AI_ROWS, max=MAX_AI_ROWS)), 400

    demo_attributes, attr_err = _validate_demo_attributes(payload.get("demo_attributes"), lang)
    if attr_err:
        return jsonify(error=attr_err), 400
    demo_attr_columns = _demo_attr_columns(demo_attributes)
    categorical_attrs_by_col = {a["column"]: a for a in demo_attributes if a["type"] == "categorical"}

    codebook, _cb_err = _validate_codebook(payload.get("codebook"), lang)
    if codebook:
        _all_qual = set(_split_codebook_columns(codebook)[1])
        qual_columns = [c for c in columns if c in _all_qual]
    else:
        qual_columns = []
    qual_col_set = set(qual_columns)
    likert_columns = [c for c in columns if c not in qual_col_set]

    all_columns = [PERSONA_COLUMN] + list(columns) + DEMO_COLUMNS + demo_attr_columns
    col_set = set(all_columns)
    for row in rows:
        if not isinstance(row, dict) or set(row.keys()) != col_set:
            return jsonify(error=t("err_ai_gen_finalize_shape_mismatch", lang)), 400
        for k, v in row.items():
            if k == PERSONA_COLUMN or k in qual_col_set:
                if not str(v).strip():
                    return jsonify(error=t("err_ai_gen_finalize_shape_mismatch", lang)), 400
                continue
            if k == "resp_gender":
                if str(v).strip().lower() not in GENDER_VALUES:
                    return jsonify(error=t("err_ai_gen_finalize_shape_mismatch", lang)), 400
                continue
            if k in categorical_attrs_by_col:
                if _normalize_categorical_value(v, categorical_attrs_by_col[k]["options"]) is None:
                    return jsonify(error=t("err_ai_gen_finalize_shape_mismatch", lang)), 400
                continue
            try:
                int(v)
            except (TypeError, ValueError):
                return jsonify(error=t("err_ai_gen_finalize_shape_mismatch", lang)), 400

    full_df = pd.DataFrame(rows, columns=all_columns)
    indicator_df = full_df[likert_columns].astype(int)
    demo_df = full_df[[PERSONA_COLUMN] + DEMO_COLUMNS + demo_attr_columns + qual_columns].copy()
    demo_df[PERSONA_COLUMN] = demo_df[PERSONA_COLUMN].astype(str).str.strip()
    demo_df["resp_age"] = demo_df["resp_age"].astype(int)
    demo_df["resp_gender"] = demo_df["resp_gender"].astype(str).str.strip().str.lower()
    for attr in demo_attributes:
        col = attr["column"]
        if attr["type"] == "numeric":
            demo_df[col] = demo_df[col].astype(int)
        else:
            demo_df[col] = demo_df[col].apply(lambda v, opts=attr["options"]: _normalize_categorical_value(v, opts))
    for col in qual_columns:
        demo_df[col] = demo_df[col].astype(str).str.strip()
    # Explicit join key between the SEM indicator file and the respondent
    # profile file -- without this, the only correspondence between the two
    # is "same row order", which is invisible once either file is opened,
    # sorted, or filtered independently (e.g. in Excel).
    demo_df.insert(0, "respondent_id", [f"R{i + 1:04d}" for i in range(len(demo_df))])

    file_id = uuid.uuid4().hex
    dest = os.path.join(_upload_dir(), file_id + ".csv")
    indicator_df.to_csv(dest, index=False)

    demographics = payload.get("demographics") or {}
    meta = {
        "file_id": file_id,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "codebook": codebook or [{"column": c, "question_text": ""} for c in columns],
        "demographics": demographics,
        "demo_attributes": demo_attributes,
        "provider": payload.get("provider"),
        "model": payload.get("model"),
        "temperature": payload.get("temperature"),
        "likert_scale": payload.get("likert_scale"),
        "n_rows": int(indicator_df.shape[0]),
        "batches": payload.get("batches") or [],
        "construct_theories": payload.get("construct_theories") or {},
    }
    _save_ai_gen_metadata(file_id, demo_df, meta)

    descriptive_stats = _compute_descriptive_stats(indicator_df, demo_df, demo_attributes)

    return jsonify(
        file_id=file_id,
        filename=filename,
        columns=list(indicator_df.columns),
        numeric_columns=list(indicator_df.columns),
        n_rows=int(indicator_df.shape[0]),
        preview=_clean(indicator_df.head(10).to_dict(orient="records")),
        demographics_summary=demographics,
        demo_attributes=demo_attributes,
        descriptive_stats=descriptive_stats,
    )


@ai_data_gen_api.get("/ai_data_gen/download")
def download():
    lang = get_lang({"lang": request.args.get("lang")})
    file_id = request.args.get("file_id") or ""
    filename = request.args.get("filename") or "ai_generated_survey.csv"
    matches = [p for p in os.listdir(_upload_dir()) if p.startswith(file_id)] if file_id else []
    if not matches:
        return jsonify(error=t("err_analyze_file_not_found", lang)), 404
    saved_path = os.path.join(_upload_dir(), matches[0])
    return send_file(saved_path, as_attachment=True, download_name=filename, mimetype="text/csv")


def _write_sheet(ws, rows: list[list]) -> None:
    for row in rows:
        ws.append(row)


@ai_data_gen_api.get("/ai_data_gen/export")
def export_full():
    lang = get_lang({"lang": request.args.get("lang")})
    file_id = request.args.get("file_id") or ""

    data_matches = [p for p in os.listdir(_upload_dir()) if p.startswith(file_id) and p.lower().endswith(".csv")] if file_id else []
    if not data_matches:
        return jsonify(error=t("err_analyze_file_not_found", lang)), 404
    indicator_df = pd.read_csv(os.path.join(_upload_dir(), data_matches[0]))

    meta, demo_df = _load_ai_gen_metadata(file_id)
    if meta is None:
        return jsonify(error=t("err_ai_gen_export_not_found", lang)), 404

    demo_attributes = meta.get("demo_attributes") or []
    stats = _compute_descriptive_stats(indicator_df, demo_df, demo_attributes)
    demographics = meta.get("demographics") or {}

    wb = Workbook()

    ws = wb.active
    ws.title = "Survey Data"
    survey_header = ["respondent_id"] + list(indicator_df.columns)
    survey_rows = [
        [rid] + row
        for rid, row in zip(demo_df["respondent_id"].tolist(), indicator_df.values.tolist())
    ]
    _write_sheet(ws, [survey_header] + survey_rows)

    def _attr_definition_desc(attr: dict) -> str:
        if attr["type"] == "numeric":
            return f"numeric ({attr['min']}-{attr['max']})"
        return f"categorical ({', '.join(attr['options'])})"

    ws = wb.create_sheet("Respondent Profile")
    _write_sheet(ws, [
        ["Occupation / education", demographics.get("occupation") or ""],
        ["Location", demographics.get("location") or ""],
        ["Target population", demographics.get("target_population") or ""],
        ["Declared age range", f"{demographics.get('age_min') or ''}-{demographics.get('age_max') or ''}"],
        ["Declared gender mix", demographics.get("gender_mix") or ""],
    ] + [
        [f"Additional attribute: {attr['name']}", _attr_definition_desc(attr)] for attr in demo_attributes
    ] + [
        [],
        list(demo_df.columns),
    ] + demo_df.values.tolist())

    ws = wb.create_sheet("Descriptive Statistics")
    _write_sheet(ws, [["Indicator", "Mean", "Std Dev", "Min", "Max"]] + [
        [col, s["mean"], s["std"], s["min"], s["max"]] for col, s in stats["indicators"].items()
    ])
    ws.append([])
    ws.append(["Age", "Mean", "Std Dev", "Min", "Max"])
    age = stats["demographics"]["age"]
    ws.append(["resp_age", age["mean"], age["std"], age["min"], age["max"]])
    ws.append([])
    ws.append(["Gender", "Count", "Percent"])
    for gender, g_stats in stats["demographics"]["gender"].items():
        ws.append([gender, g_stats["count"], g_stats["pct"]])
    for attr_stats in stats["demographics"]["custom"]:
        ws.append([])
        if attr_stats["type"] == "numeric":
            ws.append([attr_stats["name"], "Mean", "Std Dev", "Min", "Max"])
            s = attr_stats["stats"]
            ws.append([attr_stats["column"], s["mean"], s["std"], s["min"], s["max"]])
        else:
            ws.append([attr_stats["name"], "Count", "Percent"])
            for option, o_stats in attr_stats["counts"].items():
                ws.append([option, o_stats["count"], o_stats["pct"]])

    ws = wb.create_sheet("Prompt Transparency")
    _write_sheet(ws, [["Rows", "Provider", "Model", "Temperature", "System Prompt", "User Prompt"]])
    for batch in meta.get("batches") or []:
        ws.append([
            f"{batch.get('start_row')}-{batch.get('end_row')}",
            meta.get("provider"),
            meta.get("model"),
            meta.get("temperature"),
            batch.get("system_prompt"),
            batch.get("user_prompt"),
        ])

    construct_theories = meta.get("construct_theories") or {}
    if construct_theories:
        ws = wb.create_sheet("References")
        _write_sheet(ws, [["Construct", "APA Citation", "DOI"]])
        for construct_name, theory in construct_theories.items():
            doi = (theory or {}).get("doi") or ""
            ws.append([construct_name, (theory or {}).get("citation_apa") or "", f"https://doi.org/{doi}" if doi else ""])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="ai_generated_survey_full.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
